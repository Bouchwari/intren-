"""Read from and write to Excel files for the cafeteria system."""
from pathlib import Path
import sys
from typing import List, Optional, Sequence

import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation

from config.settings import (
    GRANT_FULL, GRANT_HALF,
    SECTION_CANTINE, SECTION_DAR_TALIB, SECTION_INTERNAT,
    SECTION_LABELS,
)
from core.models import LevelOption, Student


_LEVEL_OPTIONS_CACHE: list[str] | None = None
_LEVEL_CATALOG_CACHE: list[LevelOption] | None = None

_CYCLE_LABELS = {
    "1A": "الابتدائي",
    "2A": "الإعدادي",
    "3A": "التأهيلي",
    "4A": "التقني العالي",
    "5A": "التقني العالي",
}


# Common Arabic column header variants exported by Massar and similar systems.
# Each list is tried in order; first match wins.
_NAME_VARIANTS: List[str] = [
    "الاسم الشخصي والعائلي", "الاسم الكامل", "اسم التلميذ", "اسم التلميذ(ة)",
    "الاسم واللقب", "الاسم العائلي والشخصي", "الاسم",
]
_SURNAME_VARIANTS: List[str] = [
    "النسب", "اللقب", "الاسم العائلي",
]
_MASSAR_VARIANTS: List[str] = [
    "رقم مسار", "رمز مسار", "مسار", "رقم التلميذ",
]
_CLASS_VARIANTS: List[str] = [
    "الفصل", "القسم", "المستوى والفصل", "المستوى",
]
_CYCLE_VARIANTS: List[str] = [
    "السلك", "رمز السلك",
]
_EDUCATION_TYPE_VARIANTS: List[str] = [
    "نوع التعليم",
]
_GENDER_VARIANTS: List[str] = [
    "الجنس", "النوع",
]
_BIRTH_DATE_VARIANTS: List[str] = [
    "تاريخ الازدياد", "تاريخ الميلاد", "تاريخ الولادة",
]
_BIRTH_PLACE_VARIANTS: List[str] = [
    "مكان الازدياد", "مكان الميلاد", "مكان الولادة",
]
_GRANT_VARIANTS: List[str] = [
    "رقم المنحة", "رقم الاستفادة", "رقم وثيقة الاستفادة",
]
_GRANT_TYPE_VARIANTS: List[str] = [
    "نوع المنحة",
]
_SECTION_VARIANTS: List[str] = [
    "بنية الاستقبال", "بنية الإستقبال", "مكان الإيواء", "مكان الايواء", "الاستقبال", "القطاع",
]


def _template_candidates(name: str) -> list[Path]:
    """Return likely template locations for dev runs and PyInstaller builds."""
    candidates = [
        Path.cwd() / "templets" / name,
        Path(__file__).resolve().parents[2] / "templets" / name,
    ]
    if getattr(sys, "frozen", False):
        candidates.extend([
            Path(sys.executable).resolve().parent / "templets" / name,
            Path(getattr(sys, "_MEIPASS", "")).resolve() / "templets" / name,
        ])
    return candidates


def _find_template_file(name: str) -> Path | None:
    for candidate in _template_candidates(name):
        if candidate.exists():
            return candidate
    return None


def load_level_catalog() -> list[LevelOption]:
    """Load level catalog from the helper workbook, grouped for UI filtering."""
    global _LEVEL_CATALOG_CACHE
    if _LEVEL_CATALOG_CACHE is not None:
        return list(_LEVEL_CATALOG_CACHE)

    path = _find_template_file("Liste_internes.xlsx")
    if path is None:
        _LEVEL_CATALOG_CACHE = []
        return []

    catalog: list[LevelOption] = []
    seen: set[str] = set()
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        sheet = wb["niv"] if "niv" in wb.sheetnames else wb.active
        for row in sheet.iter_rows(values_only=True):
            if len(row) < 5:
                continue
            level = str(row[1] or "").strip()
            cycle_code = str(row[3] or "").strip()
            education_type = str(row[4] or "").strip()
            key = f"{cycle_code}|{education_type}|{level}"
            if not level or key in seen:
                continue
            seen.add(key)
            catalog.append(
                LevelOption(
                    cycle_code=cycle_code,
                    cycle_label=_CYCLE_LABELS.get(cycle_code, cycle_code),
                    education_type=education_type,
                    level_name=level,
                )
            )
        wb.close()
    except Exception:
        catalog = []

    _LEVEL_CATALOG_CACHE = catalog
    return list(_LEVEL_CATALOG_CACHE)


def load_level_options() -> list[str]:
    """Load level names from the helper workbook, preserving current callers."""
    global _LEVEL_OPTIONS_CACHE
    if _LEVEL_OPTIONS_CACHE is not None:
        return list(_LEVEL_OPTIONS_CACHE)

    levels: list[str] = []
    seen: set[str] = set()
    for option in load_level_catalog():
        if option.level_name in seen:
            continue
        seen.add(option.level_name)
        levels.append(option.level_name)
    _LEVEL_OPTIONS_CACHE = levels
    return list(_LEVEL_OPTIONS_CACHE)


def _unique(values: Sequence[str]) -> list[str]:
    seen: set[str] = set()
    result: list[str] = []
    for value in values:
        clean = str(value or "").strip()
        if not clean or clean in seen:
            continue
        seen.add(clean)
        result.append(clean)
    return result


def infer_level_parts(level_name: str) -> tuple[str, str]:
    """Return cycle label and education type for a level name when known."""
    normalized_level = _normalize(level_name)
    for option in load_level_catalog():
        if _normalize(option.level_name) == normalized_level:
            return option.cycle_label, option.education_type
    return "", ""


def _normalize(value: object) -> str:
    """Normalize Arabic-ish headers while keeping the original workbook flexible."""
    text = str(value or "").strip()
    text = text.replace("\n", " ").replace("\r", " ").replace("\t", " ")
    text = " ".join(text.split())
    for src, dst in (("أ", "ا"), ("إ", "ا"), ("آ", "ا"), ("ى", "ي")):
        text = text.replace(src, dst)
    for token in ("(*)", "*", ":", "："):
        text = text.replace(token, "")
    return text.strip().lower()


def _find_column(headers: List[str], variants: List[str]) -> Optional[int]:
    """Return 0-based column index of the first matching header, or None."""
    normalized_variants = [_normalize(v) for v in variants]
    for variant in variants:
        normalized_variant = _normalize(variant)
        for i, header in enumerate(headers):
            normalized_header = _normalize(header)
            if not normalized_header:
                continue
            if (
                normalized_variant == normalized_header
                or normalized_variant in normalized_header
                or normalized_header in normalized_variant
            ):
                return i
    for i, header in enumerate(headers):
        normalized_header = _normalize(header)
        if not normalized_header:
            continue
        if any(v in normalized_header for v in normalized_variants if v):
            return i
    return None


def _cell_str(row: tuple, idx: Optional[int]) -> str:
    """Safely extract a string value from a row by column index."""
    if idx is None or idx >= len(row):
        return ""
    val = row[idx]
    return str(val).strip() if val is not None else ""


def _gender_key(raw: str) -> str:
    text = _normalize(raw)
    if not text:
        return ""
    if "انث" in text or text in {"f", "female", "femme", "feminin"}:
        return "female"
    if "ذكر" in text or text in {"m", "male", "homme", "masculin"}:
        return "male"
    return raw.strip()


def _gender_label(key: str) -> str:
    return {"male": "ذكر", "female": "أنثى"}.get(key, key)


def _section_key(raw: str) -> str:
    text = _normalize(raw)
    if not text:
        return SECTION_CANTINE
    if "دار الطالب" in text or "دار الطالبه" in text:
        return SECTION_DAR_TALIB
    if "داخلي" in text or "داخليه" in text or "القسم الداخلي" in text:
        return SECTION_INTERNAT
    if "مطعم" in text:
        return SECTION_CANTINE
    return raw.strip()


def _cycle_label(raw: str) -> str:
    text = str(raw or "").strip()
    if not text:
        return ""
    if text in _CYCLE_LABELS:
        return _CYCLE_LABELS[text]
    normalized = _normalize(text)
    for label in _CYCLE_LABELS.values():
        if _normalize(label) == normalized:
            return label
    if "ابتدائي" in normalized:
        return "الابتدائي"
    if "اعدادي" in normalized:
        return "الإعدادي"
    if "تاهيلي" in normalized or "باكالوريا" in normalized or "جذع" in normalized:
        return "التأهيلي"
    if "تقني" in normalized:
        return "التقني العالي"
    return text


def _grant_key(raw: str) -> str:
    text = _normalize(raw)
    if not text:
        return GRANT_FULL
    if "نصف" in text or "demi" in text or "half" in text:
        return GRANT_HALF
    if "كامل" in text or "complete" in text or "full" in text:
        return GRANT_FULL
    return raw.strip()


def _grant_label(key: str) -> str:
    return {GRANT_FULL: "منحة كاملة", GRANT_HALF: "نصف منحة"}.get(key, key)


def _score_header_row(headers: List[str]) -> int:
    score = 0
    has_name = _find_column(headers, _NAME_VARIANTS) is not None
    has_surname = _find_column(headers, _SURNAME_VARIANTS) is not None
    if has_name or has_surname:
        score += 3
    for variants in (
        _MASSAR_VARIANTS, _CLASS_VARIANTS, _GENDER_VARIANTS,
        _CYCLE_VARIANTS, _EDUCATION_TYPE_VARIANTS,
        _BIRTH_DATE_VARIANTS, _BIRTH_PLACE_VARIANTS,
        _GRANT_VARIANTS, _GRANT_TYPE_VARIANTS, _SECTION_VARIANTS,
    ):
        if _find_column(headers, variants) is not None:
            score += 1
    return score


def _locate_header(all_rows: list[tuple]) -> tuple[int, List[str]]:
    best_row = -1
    best_headers: List[str] = []
    best_score = 0
    for row_index, row in enumerate(all_rows[:40]):
        headers = [str(cell).strip() if cell is not None else "" for cell in row]
        score = _score_header_row(headers)
        if score > best_score:
            best_row = row_index
            best_headers = headers
            best_score = score
    if best_score < 3:
        return -1, []
    return best_row, best_headers


def _find_import_source(wb: openpyxl.Workbook) -> tuple[list[tuple], int, List[str]]:
    best_rows: list[tuple] = []
    best_header_row = -1
    best_headers: List[str] = []
    best_score = 0

    for sheet in wb.worksheets:
        rows = list(sheet.iter_rows(values_only=True))
        header_row, headers = _locate_header(rows)
        if header_row < 0:
            continue
        score = _score_header_row(headers)
        if score > best_score:
            best_rows = rows
            best_header_row = header_row
            best_headers = headers
            best_score = score

    return best_rows, best_header_row, best_headers


def read_students_from_excel(file_path: Path) -> List[Student]:
    """
    Load students from a Massar-format (or template-format) Excel file.
    Auto-detects column positions from the header row.
    Raises ValueError with an Arabic message if the file is unreadable.
    """
    try:
        wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError(f"تعذر فتح الملف: {exc}") from exc

    all_rows, header_row, headers = _find_import_source(wb)
    wb.close()

    if len(all_rows) < 2 or header_row < 0:
        raise ValueError("الملف فارغ أو لا يحتوي على بيانات.")

    col_name        = _find_column(headers, _NAME_VARIANTS)
    col_surname     = _find_column(headers, _SURNAME_VARIANTS)
    col_massar      = _find_column(headers, _MASSAR_VARIANTS)
    col_cycle       = _find_column(headers, _CYCLE_VARIANTS)
    col_edu_type    = _find_column(headers, _EDUCATION_TYPE_VARIANTS)
    col_class       = _find_column(headers, _CLASS_VARIANTS)
    col_gender      = _find_column(headers, _GENDER_VARIANTS)
    col_birth_date  = _find_column(headers, _BIRTH_DATE_VARIANTS)
    col_birth_place = _find_column(headers, _BIRTH_PLACE_VARIANTS)
    col_grant       = _find_column(headers, _GRANT_VARIANTS)
    col_grant_type  = _find_column(headers, _GRANT_TYPE_VARIANTS)
    col_section     = _find_column(headers, _SECTION_VARIANTS)

    if col_name is None and col_surname is None:
        raise ValueError(
            "لم يتم العثور على عمود اسم التلميذ في الملف.\n"
            "تأكد من أن الملف يحتوي على عمود بعنوان 'الاسم الكامل' أو 'الاسم الشخصي والعائلي'.\n"
            "يمكنك تحميل النموذج الجاهز من داخل البرنامج."
        )

    students: List[Student] = []
    for row in all_rows[header_row + 1:]:
        name = _cell_str(row, col_name)
        surname = _cell_str(row, col_surname)
        name_header = _normalize(headers[col_name]) if col_name is not None and col_name < len(headers) else ""
        if surname and name and name_header in {_normalize("الاسم"), _normalize("الاسم الشخصي")}:
            name = f"{surname} {name}".strip()
        elif surname and not name:
            name = surname

        if not name:
            continue
        if _normalize(name) in {_normalize(v) for v in _NAME_VARIANTS + _SURNAME_VARIANTS}:
            continue

        student_class = _cell_str(row, col_class)
        cycle = _cycle_label(_cell_str(row, col_cycle))
        education_type = _cell_str(row, col_edu_type)
        inferred_cycle, inferred_type = infer_level_parts(student_class)
        cycle = cycle or inferred_cycle
        education_type = education_type or inferred_type

        students.append(
            Student(
                full_name=name,
                massar_number=_cell_str(row, col_massar),
                gender=_gender_key(_cell_str(row, col_gender)),
                cycle=cycle,
                education_type=education_type,
                student_class=student_class,
                birth_date=_cell_str(row, col_birth_date),
                birth_place=_cell_str(row, col_birth_place),
                grant_number=_cell_str(row, col_grant),
                section=_section_key(_cell_str(row, col_section)),
                grant_type=_grant_key(_cell_str(row, col_grant_type)),
            )
        )

    if not students:
        raise ValueError("لم يتم العثور على تلاميذ في الملف.")

    return students


def write_students_template(
    file_path: Path,
    level_catalog: Optional[Sequence[LevelOption]] = None,
) -> None:
    """Create a blank Excel template the user can fill in and import."""
    wb = openpyxl.Workbook()
    sheet = wb.active
    sheet.title = "التلاميذ"
    sheet.sheet_view.rightToLeft = True
    sheet.append([
        "ر.ت", "الاسم الشخصي والعائلي للتلميذ(ة)", "رقم مسار", "الجنس",
        "رقم المنحة", "السلك", "نوع التعليم", "المستوى", "بنية الاستقبال",
        "تاريخ الازدياد", "مكان الازدياد",
    ])
    # One example row so the user sees the expected format
    sheet.append([
        1, "محمد الأمين العمراني", "H123456789", "ذكر",
        "12345", "الإعدادي", "عام", "الأولى إعدادي عام", "القسم الداخلي",
        "2010-05-15", "الرباط",
    ])
    gender_validation = DataValidation(type="list", formula1='"ذكر,أنثى"', allow_blank=True)
    section_values = ",".join(SECTION_LABELS.values())
    section_validation = DataValidation(type="list", formula1=f'"{section_values}"', allow_blank=True)
    sheet.add_data_validation(gender_validation)
    sheet.add_data_validation(section_validation)
    gender_validation.add("D2:D1000")
    section_validation.add("I2:I1000")

    catalog = list(level_catalog) if level_catalog is not None else load_level_catalog()
    if catalog:
        choices_sheet = wb.create_sheet("_choices")
        cycle_choices = _unique([option.cycle_label for option in catalog])
        type_choices = _unique([option.education_type for option in catalog])
        for row_index, value in enumerate(cycle_choices, start=1):
            choices_sheet.cell(row_index, 1, value)
        for row_index, value in enumerate(type_choices, start=1):
            choices_sheet.cell(row_index, 2, value)
        choices_sheet.sheet_state = "hidden"

        if cycle_choices:
            cycle_validation = DataValidation(
                type="list",
                formula1=f"'_choices'!$A$1:$A${len(cycle_choices)}",
                allow_blank=True,
            )
            sheet.add_data_validation(cycle_validation)
            cycle_validation.add("F2:F1000")

        if type_choices:
            type_validation = DataValidation(
                type="list",
                formula1=f"'_choices'!$B$1:$B${len(type_choices)}",
                allow_blank=True,
            )
            sheet.add_data_validation(type_validation)
            type_validation.add("G2:G1000")

        levels_sheet = wb.create_sheet("_levels")
        for row_index, option in enumerate(catalog, start=1):
            levels_sheet.cell(row_index, 1, option.cycle_label)
            levels_sheet.cell(row_index, 2, option.education_type)
            levels_sheet.cell(row_index, 3, option.level_name)
        levels_sheet.sheet_state = "hidden"
        level_validation = DataValidation(
            type="list",
            formula1=f"'_levels'!$C$1:$C${len(catalog)}",
            allow_blank=True,
        )
        sheet.add_data_validation(level_validation)
        level_validation.add("H2:H1000")
    wb.save(file_path)


def write_students_to_excel(students: List[Student], file_path: Path) -> None:
    """Export all students to an Excel file."""
    wb = openpyxl.Workbook()
    sheet = wb.active
    sheet.title = "التلاميذ"
    sheet.sheet_view.rightToLeft = True
    sheet.append([
        "ر.ت", "الاسم الشخصي والعائلي للتلميذ(ة)", "رقم مسار", "الجنس",
        "رقم المنحة", "السلك", "نوع التعليم", "المستوى", "بنية الاستقبال",
        "تاريخ الازدياد", "مكان الازدياد", "نوع المنحة",
    ])
    for index, s in enumerate(students, start=1):
        sheet.append([
            index, s.full_name, s.massar_number, _gender_label(s.gender),
            s.grant_number, s.cycle, s.education_type, s.student_class, SECTION_LABELS.get(s.section, s.section),
            s.birth_date, s.birth_place, _grant_label(s.grant_type),
        ])
    wb.save(file_path)
