"""
src/ui/expense_statement_screen.py
Expense statement (بيان المصاريف) — detailed financial breakdown of meal costs
by beneficiary category (ممنوح / مؤد / متمم) per meal per month.
Includes summary cards, detailed table, total cost banner, and Excel export.
"""
import datetime
import logging
import os
from pathlib import Path
from typing import Dict, List, Optional

from PySide6.QtCore import Qt
from PySide6.QtGui import QColor, QFont
from PySide6.QtWidgets import (
    QComboBox, QFileDialog, QFrame, QGroupBox, QHBoxLayout,
    QHeaderView, QLabel, QMessageBox, QPushButton,
    QScrollArea, QTableWidget, QTableWidgetItem, QVBoxLayout, QWidget,
)

from config.settings import (
    COLOR_ACCENT, COLOR_BORDER, COLOR_DANGER, COLOR_SUCCESS,
    COLOR_SURFACE, COLOR_TEXT_PRIMARY, COLOR_TEXT_SECONDARY,
    MEAL_FTOUR, MEAL_GHADA, MEAL_ASHA, MEAL_LABELS,
)
from data.database import (
    get_expense_data, get_months_with_data, get_school_settings,
)

# ── Arabic strings ────────────────────────────────────────────────────────────
_TITLE       = "بيان المصاريف"
_SUBTITLE    = "البيان التفصيلي لمصاريف الإطعام المدرسي الشهري"
_BTN_GEN     = "🔄  توليد البيان"
_BTN_EXCEL   = "📊  تصدير إلى Excel"
_LBL_MONTH   = "الشهر:"
_LBL_YEAR    = "السنة:"
_NO_DATA     = "لا توجد بيانات لهذا الشهر.\nأدخل بيانات ورقة الاتصال أولاً."
_EXCEL_SAVED = "تم تصدير البيان بنجاح إلى:\n"

_ARABIC_MONTHS = [
    "يناير", "فبراير", "مارس", "أبريل", "مايو", "يونيو",
    "يوليوز", "غشت", "شتنبر", "أكتوبر", "نونبر", "دجنبر",
]

# Table columns — 11 total
_COLS = [
    "الوجبة",
    "إع. ممنوح", "إع. مؤد", "إع. متمم",
    "تأ. ممنوح", "تأ. مؤد", "تأ. متمم",
    "معلمون", "المجموع",
    "ثمن الوجبة", "المبلغ (د.م)",
]

# Category summing keys for summary cards
_GRANT_KEYS   = ("cg", "qg")   # ممنوح
_PAYING_KEYS  = ("cp", "qp")   # مؤد
_COMPL_KEYS   = ("cc", "qc")   # متمم

_MEAL_COLORS = {
    MEAL_FTOUR: "#f59e0b",
    MEAL_GHADA: COLOR_ACCENT,
    MEAL_ASHA:  "#7c3aed",
}


def _titem(text: str, bold: bool = False, bg: str = "", fg: str = "",
           align: Qt.AlignmentFlag = Qt.AlignmentFlag.AlignCenter) -> QTableWidgetItem:
    item = QTableWidgetItem(text)
    item.setTextAlignment(int(align | Qt.AlignmentFlag.AlignVCenter))
    item.setFlags(Qt.ItemFlag.ItemIsEnabled | Qt.ItemFlag.ItemIsSelectable)
    f = QFont(); f.setBold(bold); item.setFont(f)
    if bg: item.setBackground(QColor(bg))
    if fg: item.setForeground(QColor(fg))
    return item


_LOGGER = logging.getLogger(__name__)


def _safe_price(value: object) -> float:
    """Return a numeric meal price without crashing on invalid saved text."""
    try:
        return float(value or 0)
    except (TypeError, ValueError):
        return 0.0


class ExpenseStatementScreen(QWidget):

    def __init__(self) -> None:
        super().__init__()
        self.setStyleSheet(f"background:{COLOR_SURFACE};")
        self._expense_data: List[dict] = []
        self._prices: dict = {}
        self._month_str: str = ""
        self._build_ui()
        now = datetime.date.today()
        self._month_combo.setCurrentIndex(now.month - 1)
        self._year_combo.setCurrentText(str(now.year))

    # ── Build ──────────────────────────────────────────────────────────────

    def _build_ui(self) -> None:
        root = QVBoxLayout(self)
        root.setContentsMargins(0, 0, 0, 0)

        scroll = QScrollArea()
        scroll.setWidgetResizable(True)
        scroll.setFrameShape(QFrame.Shape.NoFrame)

        content = QWidget()
        content.setStyleSheet("background:transparent;")
        inner = QVBoxLayout(content)
        inner.setContentsMargins(28, 22, 28, 28)
        inner.setSpacing(18)

        inner.addLayout(self._build_header())
        inner.addLayout(self._build_selector_bar())
        inner.addLayout(self._build_summary_cards())
        inner.addWidget(self._build_table_card())
        inner.addWidget(self._build_cost_banner())
        inner.addStretch()

        scroll.setWidget(content)
        root.addWidget(scroll)

    def _build_header(self) -> QVBoxLayout:
        col = QVBoxLayout()
        title = QLabel(_TITLE)
        f = QFont(); f.setPointSize(17); f.setBold(True); title.setFont(f)
        title.setStyleSheet(f"color:{COLOR_TEXT_PRIMARY};")
        sub = QLabel(_SUBTITLE)
        sub.setStyleSheet(f"color:{COLOR_TEXT_SECONDARY}; font-size:12px;")
        col.addWidget(title); col.addWidget(sub)
        return col

    def _build_selector_bar(self) -> QHBoxLayout:
        row = QHBoxLayout(); row.setSpacing(8)

        def lbl(t): return QLabel(t, styleSheet=f"font-size:13px; color:{COLOR_TEXT_PRIMARY};")

        row.addWidget(lbl(_LBL_MONTH))
        self._month_combo = QComboBox()
        self._month_combo.setMinimumHeight(36); self._month_combo.setMinimumWidth(130)
        for m in _ARABIC_MONTHS:
            self._month_combo.addItem(m)
        self._month_combo.setStyleSheet(
            f"border:1px solid {COLOR_BORDER}; border-radius:6px; padding:4px 8px; font-size:13px;")
        row.addWidget(self._month_combo)

        row.addWidget(lbl(_LBL_YEAR))
        self._year_combo = QComboBox()
        self._year_combo.setMinimumHeight(36); self._year_combo.setMinimumWidth(90)
        cur_year = datetime.date.today().year
        for y in range(cur_year + 1, cur_year - 5, -1):
            self._year_combo.addItem(str(y))
        self._year_combo.setStyleSheet(
            f"border:1px solid {COLOR_BORDER}; border-radius:6px; padding:4px 8px; font-size:13px;")
        row.addWidget(self._year_combo)

        # Quick jump
        self._quick = QComboBox()
        self._quick.setMinimumHeight(36); self._quick.setMinimumWidth(160)
        self._quick.setPlaceholderText("الأشهر التي لها بيانات")
        self._quick.setStyleSheet(
            f"border:1px solid {COLOR_BORDER}; border-radius:6px; padding:4px 8px; font-size:13px;")
        self._quick.currentTextChanged.connect(self._on_quick_jump)
        self._refresh_quick()
        row.addWidget(self._quick)
        row.addStretch()

        gen_btn = QPushButton(_BTN_GEN)
        gen_btn.setMinimumHeight(38)
        gen_btn.setStyleSheet(
            f"background:{COLOR_ACCENT}; color:white; border-radius:7px;"
            "padding:0 18px; font-size:13px; font-weight:bold;")
        gen_btn.clicked.connect(self._generate)
        row.addWidget(gen_btn)

        xls_btn = QPushButton(_BTN_EXCEL)
        xls_btn.setMinimumHeight(38)
        xls_btn.setStyleSheet(
            f"background:#16a34a; color:white; border-radius:7px;"
            "padding:0 14px; font-size:13px; font-weight:bold;")
        xls_btn.clicked.connect(self._export_excel)
        row.addWidget(xls_btn)

        return row

    def _build_summary_cards(self) -> QHBoxLayout:
        """5 summary cards: ممنوح | مؤد | متمم | معلمون | Grand total cost."""
        row = QHBoxLayout(); row.setSpacing(12)
        defs = [
            ("ممنوح", "#0891b2",  "grant_total"),
            ("مؤد",   "#7c3aed",  "pay_total"),
            ("متمم",  "#f59e0b",  "comp_total"),
            ("معلمون","#db2777",  "mon_total"),
            ("الإجمالي المالي", COLOR_ACCENT, "cost_total"),
        ]
        self._summary_lbls: Dict[str, QLabel] = {}
        for label, color, key in defs:
            card = QFrame()
            card.setStyleSheet(
                f"background:white; border:1px solid {COLOR_BORDER};"
                "border-radius:10px; padding:2px;"
            )
            v = QVBoxLayout(card); v.setSpacing(4); v.setContentsMargins(14, 12, 14, 12)
            lbl = QLabel(label)
            lbl.setAlignment(Qt.AlignmentFlag.AlignCenter)
            lbl.setStyleSheet(f"color:{COLOR_TEXT_SECONDARY}; font-size:11px; font-weight:bold;")
            val = QLabel("—")
            val.setAlignment(Qt.AlignmentFlag.AlignCenter)
            f = QFont(); f.setPointSize(18); f.setBold(True); val.setFont(f)
            val.setStyleSheet(f"color:{color};")
            v.addWidget(lbl); v.addWidget(val)
            self._summary_lbls[key] = val
            row.addWidget(card)
        return row

    def _build_table_card(self) -> QGroupBox:
        grp = QGroupBox("التفاصيل الشهرية")
        grp.setLayoutDirection(Qt.LayoutDirection.RightToLeft)
        grp.setStyleSheet(f"""
            QGroupBox {{
                font-size:13px; font-weight:bold; color:{COLOR_TEXT_PRIMARY};
                border:1px solid {COLOR_BORDER}; border-radius:8px;
                margin-top:10px; padding:10px;
            }}
            QGroupBox::title {{
                subcontrol-origin:margin; subcontrol-position:top right;
                padding:0 8px; right:14px;
            }}
        """)
        layout = QVBoxLayout(grp)

        # No-data label
        self._no_data_lbl = QLabel(_NO_DATA)
        self._no_data_lbl.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self._no_data_lbl.setStyleSheet(
            f"color:{COLOR_TEXT_SECONDARY}; font-size:14px; padding:40px;")
        layout.addWidget(self._no_data_lbl)

        # Expense table
        self._table = QTableWidget(0, len(_COLS))
        self._table.setHorizontalHeaderLabels(_COLS)
        self._table.setLayoutDirection(Qt.LayoutDirection.RightToLeft)
        self._table.setEditTriggers(QTableWidget.EditTrigger.NoEditTriggers)
        self._table.setSelectionMode(QTableWidget.SelectionMode.NoSelection)
        self._table.verticalHeader().setVisible(False)
        self._table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.ResizeToContents)
        self._table.horizontalHeader().setStretchLastSection(True)
        self._table.setAlternatingRowColors(True)
        self._table.setShowGrid(True)
        self._table.setStyleSheet(f"""
            QTableWidget {{
                border:1px solid {COLOR_BORDER}; border-radius:8px;
                font-size:12px; background:white;
                alternate-background-color:#f8fafc;
                gridline-color:#e2e8f0;
            }}
            QHeaderView::section {{
                background:{COLOR_ACCENT}; color:white;
                padding:8px 8px; border:none;
                font-weight:bold; font-size:11px;
            }}
            QTableWidget::item {{ padding:6px 8px; }}
        """)
        self._table.setVisible(False)
        layout.addWidget(self._table)
        return grp

    def _build_cost_banner(self) -> QFrame:
        self._banner = QFrame()
        self._banner.setVisible(False)
        self._banner.setStyleSheet(
            f"background:qlineargradient(x1:0,y1:0,x2:1,y2:0,"
            f"stop:0 #0f172a,stop:1 #1e3a5f); border-radius:10px;")
        row = QHBoxLayout(self._banner)
        row.setContentsMargins(24, 14, 24, 14)

        self._banner_meals = QLabel("—")
        self._banner_cost  = QLabel("—")
        for lbl, size in ((self._banner_meals, 20), (self._banner_cost, 22)):
            f = QFont(); f.setPointSize(size); f.setBold(True); lbl.setFont(f)

        left = QVBoxLayout()
        lbl1 = QLabel("إجمالي الوجبات المُحتسبة")
        lbl1.setStyleSheet("color:rgba(255,255,255,0.7); font-size:12px;")
        self._banner_meals.setStyleSheet("color:white;")
        left.addWidget(lbl1); left.addWidget(self._banner_meals)

        right = QVBoxLayout(); right.setAlignment(Qt.AlignmentFlag.AlignLeft)
        lbl2 = QLabel("إجمالي المبلغ المالي للشهر")
        lbl2.setStyleSheet("color:rgba(255,255,255,0.7); font-size:12px;")
        self._banner_cost.setStyleSheet("color:#fde68a;")   # amber
        right.addWidget(lbl2); right.addWidget(self._banner_cost)

        row.addLayout(left)
        row.addStretch()
        row.addLayout(right)
        return self._banner

    # ── Generate ───────────────────────────────────────────────────────────

    def _selected_month_str(self) -> str:
        year  = self._year_combo.currentText()
        month = str(self._month_combo.currentIndex() + 1).zfill(2)
        return f"{year}-{month}"

    def _generate(self) -> None:
        self._month_str = self._selected_month_str()
        settings = get_school_settings()
        self._prices = {}
        if settings:
            self._prices = {
                MEAL_FTOUR: _safe_price(settings.price_ftour),
                MEAL_GHADA: _safe_price(settings.price_ghada),
                MEAL_ASHA:  _safe_price(settings.price_asha),
            }

        self._expense_data = get_expense_data(self._month_str)
        has_data = any(
            sum(r[k] for k in ("cg","cp","cc","qg","qp","qc","mo")) > 0
            for r in self._expense_data
        )

        self._no_data_lbl.setVisible(not has_data)
        self._table.setVisible(has_data)
        self._banner.setVisible(has_data)

        if has_data:
            self._populate_table()
            self._update_summary()
        else:
            self._table.setRowCount(0)
            self._banner_meals.setText("0")
            self._banner_cost.setText("0.00  د.م")
            for lbl in self._summary_lbls.values():
                lbl.setText("0")
        self._refresh_quick()

    def _populate_table(self) -> None:
        meal_labels = {MEAL_FTOUR: MEAL_LABELS[MEAL_FTOUR],
                       MEAL_GHADA: MEAL_LABELS[MEAL_GHADA],
                       MEAL_ASHA:  MEAL_LABELS[MEAL_ASHA]}
        rows = len(self._expense_data) + 1  # +1 totals row
        self._table.setRowCount(rows)

        tot_cg=tot_cp=tot_cc=tot_qg=tot_qp=tot_qc=tot_mo=0
        tot_meals=0; tot_amount=0.0

        for i, r in enumerate(self._expense_data):
            mk   = r["meal_type"]
            color = _MEAL_COLORS.get(mk, COLOR_ACCENT)
            price = self._prices.get(mk, 0.0)
            total = r["cg"]+r["cp"]+r["cc"]+r["qg"]+r["qp"]+r["qc"]+r["mo"]
            amount = total * price

            self._table.setItem(i, 0,  _titem(meal_labels.get(mk, mk), bold=True, fg=color))
            self._table.setItem(i, 1,  _titem(str(r["cg"])))
            self._table.setItem(i, 2,  _titem(str(r["cp"])))
            self._table.setItem(i, 3,  _titem(str(r["cc"])))
            self._table.setItem(i, 4,  _titem(str(r["qg"])))
            self._table.setItem(i, 5,  _titem(str(r["qp"])))
            self._table.setItem(i, 6,  _titem(str(r["qc"])))
            self._table.setItem(i, 7,  _titem(str(r["mo"])))
            self._table.setItem(i, 8,  _titem(str(total), bold=True))
            self._table.setItem(i, 9,  _titem(f"{price:.2f}"))
            self._table.setItem(i, 10, _titem(f"{amount:.2f}", bold=True))

            tot_cg+=r["cg"]; tot_cp+=r["cp"]; tot_cc+=r["cc"]
            tot_qg+=r["qg"]; tot_qp+=r["qp"]; tot_qc+=r["qc"]
            tot_mo+=r["mo"]; tot_meals+=total; tot_amount+=amount

        # Grand total row
        r = rows - 1
        bg, fg = "#0f172a", "white"
        self._table.setItem(r, 0,  _titem("الإجمالي", bold=True, bg=bg, fg=fg))
        self._table.setItem(r, 1,  _titem(str(tot_cg), bold=True, bg=bg, fg=fg))
        self._table.setItem(r, 2,  _titem(str(tot_cp), bold=True, bg=bg, fg=fg))
        self._table.setItem(r, 3,  _titem(str(tot_cc), bold=True, bg=bg, fg=fg))
        self._table.setItem(r, 4,  _titem(str(tot_qg), bold=True, bg=bg, fg=fg))
        self._table.setItem(r, 5,  _titem(str(tot_qp), bold=True, bg=bg, fg=fg))
        self._table.setItem(r, 6,  _titem(str(tot_qc), bold=True, bg=bg, fg=fg))
        self._table.setItem(r, 7,  _titem(str(tot_mo), bold=True, bg=bg, fg=fg))
        self._table.setItem(r, 8,  _titem(str(tot_meals), bold=True, bg=bg, fg=fg))
        self._table.setItem(r, 9,  _titem("", bg=bg))
        self._table.setItem(r, 10, _titem(f"{tot_amount:.2f}", bold=True,
                                           bg=COLOR_ACCENT, fg="white"))

        self._table.setMaximumHeight(rows * 38 + 44)
        # Banner
        self._banner_meals.setText(str(tot_meals))
        self._banner_cost.setText(f"{tot_amount:,.2f}  د.م")

    def _update_summary(self) -> None:
        grant=pay=comp=mon=0; total_cost=0.0
        for r in self._expense_data:
            price = self._prices.get(r["meal_type"], 0.0)
            g = r["cg"]+r["qg"]; p = r["cp"]+r["qp"]
            c = r["cc"]+r["qc"]; m = r["mo"]
            grant+=g; pay+=p; comp+=c; mon+=m
            total_cost += (g+p+c+m) * price

        self._summary_lbls["grant_total"].setText(str(grant))
        self._summary_lbls["pay_total"].setText(str(pay))
        self._summary_lbls["comp_total"].setText(str(comp))
        self._summary_lbls["mon_total"].setText(str(mon))
        self._summary_lbls["cost_total"].setText(f"{total_cost:,.2f}")

    # ── Quick jump ─────────────────────────────────────────────────────────

    def _refresh_quick(self) -> None:
        self._quick.blockSignals(True)
        self._quick.clear()
        for m in get_months_with_data():
            self._quick.addItem(m)
        self._quick.blockSignals(False)

    def _on_quick_jump(self, month_str: str) -> None:
        if not month_str or len(month_str) != 7:
            return
        try:
            year, month = month_str.split("-")
            self._year_combo.setCurrentText(year)
            self._month_combo.setCurrentIndex(int(month) - 1)
            self._generate()
        except (ValueError, IndexError):
            _LOGGER.warning("Ignoring invalid quick-jump month: %s", month_str)

    # ── Excel export ───────────────────────────────────────────────────────

    def _export_excel(self) -> None:
        if not self._expense_data:
            QMessageBox.warning(self, "تنبيه", "ولّد البيان أولاً ثم قم بالتصدير.")
            return
        try:
            import openpyxl
            from openpyxl.styles import (
                Alignment, Border, Font, PatternFill, Side,
            )
        except ImportError:
            QMessageBox.critical(self, "خطأ", "مكتبة openpyxl غير مثبتة.\nثبّتها بـ:  pip install openpyxl")
            return

        default_name = f"بيان_المصاريف_{self._month_str}.xlsx"
        path, _ = QFileDialog.getSaveFileName(
            self, "حفظ البيان",
            str(Path.home() / default_name),
            "Excel Files (*.xlsx)"
        )
        if not path:
            return

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "بيان المصاريف"

        # Styles
        hdr_fill  = PatternFill("solid", fgColor="1e40af")  # blue-800
        tot_fill  = PatternFill("solid", fgColor="0f172a")  # slate-900
        amt_fill  = PatternFill("solid", fgColor="2563eb")  # blue-600
        hdr_font  = Font(name="Arial", bold=True, color="FFFFFF", size=11)
        tot_font  = Font(name="Arial", bold=True, color="FFFFFF", size=11)
        body_font = Font(name="Arial", size=11)
        center    = Alignment(horizontal="center", vertical="center",
                              wrap_text=True, readingOrder=2)
        thin_side = Side(style="thin", color="94a3b8")
        border    = Border(left=thin_side, right=thin_side,
                           top=thin_side, bottom=thin_side)

        # Title
        ws.merge_cells("A1:K1")
        title_cell = ws["A1"]
        title_cell.value = f"بيان المصاريف — {_ARABIC_MONTHS[self._month_combo.currentIndex()]} {self._year_combo.currentText()}"
        title_cell.font  = Font(name="Arial", bold=True, size=14, color="1e40af")
        title_cell.alignment = center
        ws.row_dimensions[1].height = 28

        settings = get_school_settings()
        ws.merge_cells("A2:K2")
        ws["A2"].value = settings.school_name if settings else ""
        ws["A2"].alignment = center
        ws["A2"].font = Font(name="Arial", size=11, color="475569")

        # Headers row 3
        ws.row_dimensions[3].height = 30
        for col_idx, hdr_text in enumerate(_COLS, 1):
            cell = ws.cell(row=3, column=col_idx, value=hdr_text)
            cell.fill      = hdr_fill
            cell.font      = hdr_font
            cell.alignment = center
            cell.border    = border

        # Data rows
        meal_labels = {MEAL_FTOUR: MEAL_LABELS[MEAL_FTOUR],
                       MEAL_GHADA: MEAL_LABELS[MEAL_GHADA],
                       MEAL_ASHA:  MEAL_LABELS[MEAL_ASHA]}
        row_num = 4
        tot_cg=tot_cp=tot_cc=tot_qg=tot_qp=tot_qc=tot_mo=0
        tot_meals=0; tot_amount=0.0

        for r in self._expense_data:
            mk    = r["meal_type"]
            price = self._prices.get(mk, 0.0)
            total = r["cg"]+r["cp"]+r["cc"]+r["qg"]+r["qp"]+r["qc"]+r["mo"]
            amount = total * price
            row_data = [
                meal_labels.get(mk, mk),
                r["cg"], r["cp"], r["cc"],
                r["qg"], r["qp"], r["qc"],
                r["mo"], total, f"{price:.2f}", f"{amount:.2f}",
            ]
            for col_idx, val in enumerate(row_data, 1):
                cell = ws.cell(row=row_num, column=col_idx, value=val)
                cell.font      = body_font
                cell.alignment = center
                cell.border    = border
            ws.row_dimensions[row_num].height = 22
            row_num += 1
            tot_cg+=r["cg"]; tot_cp+=r["cp"]; tot_cc+=r["cc"]
            tot_qg+=r["qg"]; tot_qp+=r["qp"]; tot_qc+=r["qc"]
            tot_mo+=r["mo"]; tot_meals+=total; tot_amount+=amount

        # Grand total row
        total_data = [
            "الإجمالي", tot_cg, tot_cp, tot_cc,
            tot_qg, tot_qp, tot_qc,
            tot_mo, tot_meals, "", f"{tot_amount:.2f}",
        ]
        ws.row_dimensions[row_num].height = 26
        for col_idx, val in enumerate(total_data, 1):
            cell = ws.cell(row=row_num, column=col_idx, value=val)
            cell.fill      = tot_fill if col_idx < 11 else amt_fill
            cell.font      = tot_font
            cell.alignment = center
            cell.border    = border

        # Column widths
        widths = [14, 10, 10, 10, 10, 10, 10, 10, 10, 12, 16]
        for col_idx, w in enumerate(widths, 1):
            ws.column_dimensions[
                openpyxl.utils.get_column_letter(col_idx)
            ].width = w

        ws.sheet_view.rightToLeft = True

        wb.save(path)
        QMessageBox.information(self, "تم التصدير", _EXCEL_SAVED + path)
