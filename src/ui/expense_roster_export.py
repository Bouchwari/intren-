"""
src/ui/expense_roster_export.py
Writes بيانات المصاريف — the official quarterly per-student roster, matching
`templets/بيانات مصاريف يناير.فبراير مارس 2026 -.xlsx` cell-for-cell.

Structure verified by dumping the real file (merges, row heights, column
widths, formulas, the embedded ministry header image), not inferred:

  * One sheet per cycle: ابتدائي / اعدادي / تأهيلي — the template's own
    sheet names, RTL, A4 portrait.
  * Rows 1-5 hold the ministry header image (extracted from the template
    itself, so the crest is the real one rather than a redrawn copy).
  * A6 = المؤسسة, H6 = السلك, A7 = the "بيانات مصاريف : <months> <year>"
    title, and a small grant-type tally box on the right whose figures are
    live COUNTIF formulas over the student rows, exactly as the template.
  * "1- التلاميذ الممنوحين" table: ر.ت / الاسم / رقم مسار / الجنس /
    رقم المنحة / المستوى / نوع المنحة / عدد الوجبات الغذائية / بنية الاستقبال.
    ابتدائي always gets 3 meal columns (الفطور/الغذاء/العشاء). اعدادي and
    تأهيلي get the template's two extra Ramadan columns (السحور + a second
    الفطور meaning الإفطار) ONLY for a quarter that actually contains
    Ramadan — the user asked for them to stay out of the document rather
    than sit empty for the rest of the year.
  * اعدادي and تأهيلي additionally carry the "معلموا الداخلية" table and a
    المجموع (1)/(2)/العام block; ابتدائي has neither, again per the template.

Identity data (name, Massar, gender, grant number, level, grant type,
بنية الاستقبال) is filled from the real student list. **The per-student meal
counts are deliberately left EMPTY** — the app records attendance as daily
per-cycle totals, never per pupil, so there is no honest per-student number
to put there. On a document a government body pays against, a blank the user
fills in is the only defensible option; see CLAUDE.md §7.
"""
import logging
import sys
from pathlib import Path
from typing import List, Optional
from zipfile import ZipFile

from config.settings import (
    GRANT_FULL, SECTION_CANTINE, SECTION_DAR_TALIB, SECTION_INTERNAT,
)
from core.models import SchoolSettings, Student

# ── The template this file reproduces ───────────────────────────────────────
_TEMPLATE_NAME = "بيانات مصاريف يناير.فبراير مارس 2026 -.xlsx"
_TEMPLATE_IMAGE = "xl/media/image1.png"

_ARABIC_MONTHS = [
    "يناير", "فبراير", "مارس", "أبريل", "مايو", "يونيو",
    "يوليوز", "غشت", "شتنبر", "أكتوبر", "نونبر", "دجنبر",
]

_NORMAL_MEAL_HEADERS = ["الفطور", "الغذاء", "العشاء"]
# السحور + الفطور(الإفطار) — appended only for a quarter that actually
# contains Ramadan. The user asked for these two columns to stay out of the
# document until Ramadan rather than sitting empty all year.
_RAMADAN_MEAL_HEADERS = ["السحور", "الفطور"]

# The header row differs per sheet in the template — اعدادي starts one row
# lower than the other two. Reproduced rather than normalised, per the
# standing "respect the template" rule (CLAUDE.md §9).
# (sheet, السلك caption, takes Ramadan columns, has معلموا الداخلية, header row)
_CYCLE_SHEETS = [
    ("ابتدائي", "الإبتدائي", False, False, 12),
    ("اعدادي", "الثانوي الإعدادي", True, True, 13),
    ("تأهيلي", "الثانوي التأهيلي", True, True, 12),
]

# Which cycle a Student belongs on. Student.cycle holds either the Arabic
# label or the Massar code (1A/2A/3A/4A) depending on how the roster was
# imported, so both are matched.
_CYCLE_MATCH = {
    "ابتدائي": ("ابتدائ", "1A"),
    "اعدادي": ("عدادي", "إعدادي", "2A"),
    "تأهيلي": ("تأهيل", "تاهيل", "3A", "4A"),
}

_HEADER_LABELS = [
    "ر.ت", "الاسم الشخصي والعائلي للتلميذ (ة)", "رقم مسار", "الجنس",
    "رقم المنحة", "المستوى", "نوع المنحة",
]
_MEALS_GROUP_LABEL = "عدد الوجبات الغذائية"
_COUNT_LABEL = "عدد "
_STRUCTURE_LABEL = ("بنية الاستقبال(*):\n-القسم الداخلي؛ \n-دار الطالب(ة)؛\n"
                    "- مطعم مدرسي.")
_STRUCTURE_NOTE = ".يرجى تحديد بنية الاستقبال الخاصة بكل تلميذ  : (*)"
_GRANTED_SECTION = "1- التلاميذ الممنوحين:"
_MONITORS_SECTION = "معلموا  الداخلية:"

# The tally box on the right. The template counts three wordings; the user
# confirmed نصف منحة and وجبة غذاء are the same category in practice, so a
# half grant is printed as وجبة غذاء (the wording the official document
# uses) while the نصف منحة row is kept so the box matches the template.
_TALLY_ROWS = ["منحة كاملة", "وجبة غذاء", "نصف منحة"]
_GRANT_FULL_LABEL = "منحة كاملة"
_GRANT_LUNCH_LABEL = "وجبة غذاء"

_SECTION_LABELS_OFFICIAL = {
    SECTION_INTERNAT: "القسم الداخلي",
    SECTION_DAR_TALIB: "دار الطالب(ة)",
    SECTION_CANTINE: "مطعم مدرسي",
}

_GENDER_LABELS = {"male": "ذكر", "female": "أنثى"}

_LOGGER = logging.getLogger(__name__)

_FIRST_SECTION_ROW = 11      # "1- التلاميذ الممنوحين:"
_HEADER_ROW = 12             # default header row; اعدادي overrides it to 13
_FIRST_STUDENT_ROW = 15      # = header row + 3

_COL_WIDTHS = {
    "A": 6.0, "B": 18.0, "C": 13.8, "D": 9.3, "E": 11.5, "F": 11.3, "G": 9.2,
}
_ROW_HEIGHT_HEADER = 42.0
_ROW_HEIGHT_SUBHEADER = 19.5
_ROW_HEIGHT_DATA = 21.5
_ROW_HEIGHT_TOTAL = 27.0


def _signature_spans(has_monitors: bool, last_col: int) -> List[tuple]:
    """(label, first column, last column) for the signature row.

    The template's own spans are used when the sheet has its full width
    (A:C / D:F / G:K / L:N with the Ramadan columns present, A:E / F:L on
    ابتدائي); a Ramadan-free sheet is narrower, so the groups are split
    evenly across whatever columns it actually has.
    """
    if has_monitors:
        labels = ["الحارس (ة) العام للداخلية:", "مسير(ة) المصالح المالية والمادية:",
                  "رئيس المؤسسة", "المدير الإقليمي"]
        if last_col == 14:                       # the template's own layout
            return list(zip(labels, (1, 4, 7, 12), (3, 6, 11, 14)))
    else:
        labels = ["رئيس المؤسسة:", "المدير الإقليمي:"]
        if last_col == 12:                       # the template's own layout
            return list(zip(labels, (1, 6), (5, 12)))

    width = max(1, last_col // len(labels))
    spans = []
    for index, label in enumerate(labels):
        start = 1 + index * width
        end = last_col if index == len(labels) - 1 else start + width - 1
        spans.append((label, start, min(end, last_col)))
    return spans


def _template_path() -> Optional[Path]:
    """Locate the source template — only needed for its header image."""
    runtime_base = Path(getattr(sys, "_MEIPASS", Path.cwd()))
    for base in (runtime_base, Path.cwd(), Path(__file__).resolve().parents[2]):
        candidate = base / "templets" / _TEMPLATE_NAME
        if candidate.exists():
            return candidate
    return None


def _header_image_bytes() -> Optional[bytes]:
    """The real ministry crest, read straight out of the template so the
    generated roster carries the same header rather than a redrawn one."""
    path = _template_path()
    if path is None:
        return None
    try:
        with ZipFile(path) as archive:
            return archive.read(_TEMPLATE_IMAGE)
    except Exception:
        return None


_DRAWING_XML = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
    '<xdr:wsDr xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/'
    'spreadsheetDrawing" xmlns:a="http://schemas.openxmlformats.org/drawingml/'
    '2006/main"><xdr:twoCellAnchor editAs="oneCell"><xdr:from><xdr:col>1'
    '</xdr:col><xdr:colOff>95250</xdr:colOff><xdr:row>0</xdr:row><xdr:rowOff>'
    '47625</xdr:rowOff></xdr:from><xdr:to><xdr:col>{last_col}</xdr:col>'
    '<xdr:colOff>0</xdr:colOff><xdr:row>5</xdr:row><xdr:rowOff>0</xdr:rowOff>'
    '</xdr:to><xdr:pic><xdr:nvPicPr><xdr:cNvPr id="2" name="Header"/>'
    '<xdr:cNvPicPr/></xdr:nvPicPr><xdr:blipFill><a:blip xmlns:r="http://'
    'schemas.openxmlformats.org/officeDocument/2006/relationships" '
    'r:embed="rId1"/><a:stretch><a:fillRect/></a:stretch></xdr:blipFill>'
    '<xdr:spPr><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></xdr:spPr>'
    '</xdr:pic><xdr:clientData/></xdr:twoCellAnchor></xdr:wsDr>'
)

_DRAWING_RELS_XML = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
    '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/'
    'relationships"><Relationship Id="rId1" Type="http://schemas.'
    'openxmlformats.org/officeDocument/2006/relationships/image" '
    'Target="../media/image1.png"/></Relationships>'
)

_SHEET_RELS_XML = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
    '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/'
    'relationships"><Relationship Id="rIdHdr" Type="http://schemas.'
    'openxmlformats.org/officeDocument/2006/relationships/drawing" '
    'Target="../drawings/drawing{index}.xml"/></Relationships>'
)


def _embed_header_images(path: Path, sheet_last_cols: List[int],
                         image_bytes: bytes) -> None:
    """Put the ministry crest on every sheet by writing the drawing parts
    into the saved .xlsx directly.

    openpyxl's own `add_image` needs Pillow, which this project does not
    depend on (and CLAUDE.md §9 forbids adding a library without asking), so
    the picture parts are assembled by hand instead — a few small XML files
    plus the PNG, which needs no image library at all.
    """
    import re
    import shutil
    import tempfile

    handle, temp_name = tempfile.mkstemp(suffix=".xlsx")
    import os
    os.close(handle)
    temp_path = Path(temp_name)

    try:
        with ZipFile(path) as source, ZipFile(temp_path, "w") as target:
            names = source.namelist()
            sheet_parts = sorted(
                (n for n in names if re.fullmatch(r"xl/worksheets/sheet\d+\.xml", n)),
                key=lambda n: int(re.search(r"(\d+)", n.rsplit("/", 1)[-1]).group(1)),
            )
            for item in source.infolist():
                data = source.read(item.filename)
                if item.filename in sheet_parts:
                    index = sheet_parts.index(item.filename) + 1
                    text = data.decode("utf-8")
                    if "<drawing" not in text:
                        text = text.replace(
                            "</worksheet>", '<drawing r:id="rIdHdr"/></worksheet>')
                        # The sheet element may not declare the r: prefix yet.
                        if 'xmlns:r=' not in text.split(">", 2)[1]:
                            text = text.replace(
                                "<worksheet ",
                                '<worksheet xmlns:r="http://schemas.openxmlformats.org'
                                '/officeDocument/2006/relationships" ', 1)
                    data = text.encode("utf-8")
                elif item.filename == "[Content_Types].xml":
                    text = data.decode("utf-8")
                    if 'Extension="png"' not in text:
                        text = text.replace(
                            "<Default", '<Default Extension="png" '
                            'ContentType="image/png"/><Default', 1)
                    overrides = "".join(
                        f'<Override PartName="/xl/drawings/drawing{i}.xml" '
                        f'ContentType="application/vnd.openxmlformats-officedocument'
                        f'.drawing+xml"/>' for i in range(1, len(sheet_parts) + 1)
                    )
                    text = text.replace("</Types>", overrides + "</Types>")
                    data = text.encode("utf-8")
                target.writestr(item, data)

            target.writestr("xl/media/image1.png", image_bytes)
            for index, last_col in enumerate(sheet_last_cols, start=1):
                target.writestr(
                    f"xl/drawings/drawing{index}.xml",
                    _DRAWING_XML.format(last_col=last_col))
                target.writestr(
                    f"xl/drawings/_rels/drawing{index}.xml.rels", _DRAWING_RELS_XML)
                if f"xl/worksheets/_rels/sheet{index}.xml.rels" not in names:
                    target.writestr(
                        f"xl/worksheets/_rels/sheet{index}.xml.rels",
                        _SHEET_RELS_XML.format(index=index))
        shutil.move(str(temp_path), str(path))
    finally:
        if temp_path.exists():
            temp_path.unlink()


def _cycle_of(student: Student) -> str:
    """Return the sheet name a student belongs on, defaulting to اعدادي."""
    raw = f"{student.cycle} {student.education_type}".strip()
    for sheet_name, needles in _CYCLE_MATCH.items():
        if any(needle in raw for needle in needles):
            return sheet_name
    return "اعدادي"


def _grant_label(student: Student) -> str:
    return _GRANT_FULL_LABEL if student.grant_type == GRANT_FULL else _GRANT_LUNCH_LABEL


def group_students_by_cycle(students: List[Student]) -> dict:
    """Split the roster into {sheet_name: [students]}, monitors excluded."""
    grouped: dict = {sheet[0]: [] for sheet in _CYCLE_SHEETS}
    for student in students:
        if student.is_monitor:
            continue
        grouped[_cycle_of(student)].append(student)
    return grouped


def _quarter_title(months: List[tuple]) -> str:
    """"بيانات مصاريف : يناير- فبراير-مارس 2026" — the template's own
    wording, with each month carrying its year when the quarter straddles
    31 December so the document can't certify the wrong period."""
    years = {year for year, _m in months}
    if len(years) == 1:
        names = "- ".join(_ARABIC_MONTHS[m - 1] for _y, m in months)
        return f"بيانات مصاريف : {names} {months[-1][0]}"
    names = "- ".join(f"{_ARABIC_MONTHS[m - 1]} {y}" for y, m in months)
    return f"بيانات مصاريف : {names}"


def _write_cycle_sheet(
    wb, sheet_name: str, cycle_caption: str, meal_headers: List[str],
    has_monitors: bool, students: List[Student], monitors: List[Student],
    settings: Optional[SchoolSettings], months: List[tuple],
    header_row: int = _HEADER_ROW,
) -> None:
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    ws = wb.create_sheet(title=sheet_name)
    ws.sheet_view.rightToLeft = True

    meal_count = len(meal_headers)
    structure_col = 8 + meal_count          # بنية الاستقبال starts here
    last_col = structure_col + 1            # it spans two columns
    tally_label_col = structure_col
    tally_value_col = structure_col + 1

    thin = Side(style="thin", color="FF000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    centre = Alignment(horizontal="center", vertical="center", wrap_text=True,
                       readingOrder=2)
    header_fill = PatternFill("solid", fgColor="D9D9D9")

    def cell(row: int, col: int, value=None, *, bold=False, size=11,
             fill: str = "", boxed: bool = True, align=centre):
        target = ws.cell(row=row, column=col)
        if value is not None:
            target.value = value
        target.font = Font(name="Arial", bold=bold, size=size)
        target.alignment = align
        if fill:
            target.fill = PatternFill("solid", fgColor=fill)
        if boxed:
            target.border = border
        return target

    # Rows 1-5 are reserved for the ministry crest, which is written into the
    # saved file afterwards by _embed_header_images (no Pillow needed).
    for row in range(1, 6):
        ws.row_dimensions[row].height = 22.0

    # ── Identity line + quarter title ──────────────────────────────────────
    school_name = (settings.school_name if settings else "") or ""
    ws.merge_cells(start_row=6, start_column=1, end_row=6, end_column=4)
    cell(6, 1, f"المؤسسة : {school_name}", bold=True, size=12, boxed=False,
         align=Alignment(horizontal="right", vertical="center", readingOrder=2))
    ws.merge_cells(start_row=6, start_column=8, end_row=6, end_column=last_col)
    cell(6, 8, f"السلك: {cycle_caption}", bold=True, size=12, boxed=False)

    ws.merge_cells(start_row=7, start_column=1, end_row=10, end_column=structure_col - 1)
    cell(7, 1, _quarter_title(months), bold=True, size=14, boxed=False)
    ws.row_dimensions[6].height = 21.5
    ws.row_dimensions[7].height = 27.75

    # ── Grant-type tally box — live COUNTIF formulas like the template ─────
    grant_column = get_column_letter(7)  # نوع المنحة
    last_student_row = header_row + 3 + max(len(students), 1) - 1
    tally_range = f"{grant_column}{header_row + 3}:{grant_column}{last_student_row}"
    cell(7, tally_label_col, "المجموع", bold=True, size=11)
    cell(7, tally_value_col, f"=N{8}+N{9}+N{10}".replace("N", get_column_letter(tally_value_col)),
         bold=True, size=11)
    for offset, label in enumerate(_TALLY_ROWS, start=8):
        cell(offset, tally_label_col, label, size=11)
        cell(offset, tally_value_col, f'=COUNTIF({tally_range},"{label}")', size=11)

    # ── "1- التلاميذ الممنوحين" table ───────────────────────────────────────
    cell(_FIRST_SECTION_ROW, 1, _GRANTED_SECTION, bold=True, size=12, boxed=False,
         align=Alignment(horizontal="right", vertical="center", readingOrder=2))

    def write_table_header(header_row: int) -> None:
        """The template's 3-row header: labels merged down, the meal columns
        grouped under one caption, بنية الاستقبال merged down on the right."""
        ws.row_dimensions[header_row].height = _ROW_HEIGHT_HEADER
        ws.row_dimensions[header_row + 1].height = _ROW_HEIGHT_SUBHEADER
        ws.row_dimensions[header_row + 2].height = _ROW_HEIGHT_DATA
        for index, label in enumerate(_HEADER_LABELS, start=1):
            ws.merge_cells(start_row=header_row, start_column=index,
                           end_row=header_row + 2, end_column=index)
            for row in range(header_row, header_row + 3):
                cell(row, index, label if row == header_row else None,
                     bold=True, size=11, fill="D9D9D9")
        ws.merge_cells(start_row=header_row, start_column=8,
                       end_row=header_row, end_column=8 + meal_count - 1)
        for col in range(8, 8 + meal_count):
            cell(header_row, col, _MEALS_GROUP_LABEL if col == 8 else None,
                 bold=True, size=11, fill="D9D9D9")
        for offset, meal in enumerate(meal_headers):
            cell(header_row + 1, 8 + offset, meal, bold=True, size=11, fill="D9D9D9")
            cell(header_row + 2, 8 + offset, _COUNT_LABEL, size=11, fill="D9D9D9")
        ws.merge_cells(start_row=header_row, start_column=structure_col,
                       end_row=header_row + 2, end_column=last_col)
        for row in range(header_row, header_row + 3):
            for col in (structure_col, last_col):
                cell(row, col, _STRUCTURE_LABEL if (row, col) == (header_row, structure_col)
                     else None, bold=True, size=9, fill="D9D9D9")

    def write_people(first_row: int, people: List[Student]) -> int:
        """One row per real person; meal columns intentionally left blank."""
        row = first_row
        for serial, person in enumerate(people, start=1):
            ws.row_dimensions[row].height = _ROW_HEIGHT_DATA
            cell(row, 1, serial, size=11)
            cell(row, 2, person.full_name, size=11,
                 align=Alignment(horizontal="right", vertical="center",
                                 wrap_text=True, readingOrder=2))
            cell(row, 3, person.massar_number, size=11)
            cell(row, 4, _GENDER_LABELS.get(person.gender, ""), size=11)
            cell(row, 5, person.grant_number, size=11)
            cell(row, 6, person.student_class, size=11)
            cell(row, 7, _grant_label(person), size=11)
            for offset in range(meal_count):
                cell(row, 8 + offset, None, size=11)   # filled in by hand
            ws.merge_cells(start_row=row, start_column=structure_col,
                           end_row=row, end_column=last_col)
            for col in (structure_col, last_col):
                cell(row, col,
                     _SECTION_LABELS_OFFICIAL.get(person.section, "") if col == structure_col
                     else None, size=10)
            row += 1
        return row

    def write_total(row: int, label: str) -> None:
        ws.row_dimensions[row].height = _ROW_HEIGHT_TOTAL
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        cell(row, 1, label, bold=True, size=12, fill="D9D9D9")
        cell(row, 2, None, fill="D9D9D9")
        for col in range(3, last_col + 1):
            cell(row, col, None, fill="D9D9D9")

    write_table_header(header_row)
    next_row = write_people(header_row + 3, students)
    granted_total_row = next_row
    write_total(granted_total_row, "المجموع (1)" if has_monitors else "المجموع")
    next_row = granted_total_row + 1

    if has_monitors:
        cell(next_row, 1, _MONITORS_SECTION, bold=True, size=12, boxed=False,
             align=Alignment(horizontal="right", vertical="center", readingOrder=2))
        next_row += 1
        write_table_header(next_row)
        next_row += 3
        next_row = write_people(next_row, monitors)
        write_total(next_row, "المجموع (2)")
        next_row += 1
        write_total(next_row, "المجموع العام (1)+(2)")
        next_row += 1

    # ── Footer: the بنية الاستقبال note and the signature row ──────────────
    next_row += 1
    ws.merge_cells(start_row=next_row, start_column=1, end_row=next_row, end_column=5)
    cell(next_row, 1, _STRUCTURE_NOTE, size=10, boxed=False,
         align=Alignment(horizontal="right", vertical="center", readingOrder=2))

    signature_row = next_row + 2
    ws.row_dimensions[signature_row].height = 24.0
    signers = _signature_spans(has_monitors, last_col)
    for signer, start, end in signers:
        ws.merge_cells(start_row=signature_row, start_column=start,
                       end_row=signature_row, end_column=end)
        cell(signature_row, start, signer, bold=True, size=11, boxed=False)
        # Signing space underneath, framed like the template's boxes.
        ws.merge_cells(start_row=signature_row + 1, start_column=start,
                       end_row=signature_row + 3, end_column=end)
        for row in range(signature_row + 1, signature_row + 4):
            for col in range(start, end + 1):
                cell(row, col, None)

    for letter, width in _COL_WIDTHS.items():
        ws.column_dimensions[letter].width = width
    for offset in range(meal_count):
        ws.column_dimensions[get_column_letter(8 + offset)].width = 7.5
    ws.column_dimensions[get_column_letter(structure_col)].width = 11.0
    ws.column_dimensions[get_column_letter(last_col)].width = 7.7

    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = 9  # A4
    ws.page_setup.scale = 70
    ws.print_area = (f"'{sheet_name}'!$A$1:"
                     f"${get_column_letter(last_col)}${signature_row + 3}")


def write_expense_roster_excel(
    path: Path, settings: Optional[SchoolSettings], students: List[Student],
    months: List[tuple], has_ramadan: bool = False,
) -> None:
    """Write بيانات المصاريف for the given quarter — one sheet per cycle.

    `months` is the quarter as [(year, month), ...]; `students` is the full
    roster (monitors included — they are routed to their own table).
    `has_ramadan` adds the السحور/الإفطار columns; without it they are left
    out entirely rather than printed empty all year.
    """
    import openpyxl

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    grouped = group_students_by_cycle(students)
    monitors = [s for s in students if s.is_monitor]
    image_bytes = _header_image_bytes()
    last_cols: List[int] = []

    for sheet_name, caption, takes_ramadan, has_monitors, header_row in _CYCLE_SHEETS:
        meal_headers = list(_NORMAL_MEAL_HEADERS)
        if takes_ramadan and has_ramadan:
            meal_headers += _RAMADAN_MEAL_HEADERS
        # Monitors serve the whole internat, not one cycle. They are listed
        # on اعدادي only — repeating them on every sheet that has the table
        # would claim the same people twice across the dossier.
        sheet_monitors = monitors if sheet_name == "اعدادي" else []
        _write_cycle_sheet(
            wb, sheet_name, caption, meal_headers, has_monitors,
            grouped.get(sheet_name, []), sheet_monitors, settings, months,
            header_row,
        )
        last_cols.append(7 + len(meal_headers) + 2)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)

    if image_bytes:
        try:
            _embed_header_images(path, last_cols, image_bytes)
        except Exception:
            # The roster is still valid without the crest — never lose the
            # whole export over a decorative header.
            _LOGGER.exception("Could not embed the ministry header image")
