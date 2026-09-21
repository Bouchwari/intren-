"""Create and read the Excel workbook used for bulk daily entry."""
from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any, Iterator

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from config.settings import (
    ARABIC_DAY_NAMES, MEAL_LABELS, RAMADAN_MEALS, REGULAR_MEALS,
)
from core.bulk_daily_entry import validate_bulk_daily_entries
from core.models import BulkDailyEntry, DailyAbsence, DailyContact
from core.ramadan import meals_for_date
from data.database import (
    get_all_holidays, get_day_absences, get_day_contacts,
    get_ramadan_overrides, get_school_settings,
)


SHEET_TITLE = "الإدخال اليومي"
_META_SHEET = "_meta"
_SIGNATURE = "TADBIR_BULK_DAILY_V1"
_HEADER_ROW = 5
_DATA_ROW = 6
_HEADERS = ("التاريخ", "اليوم", "الوجبة", "الحضور", "الغياب")
_MEAL_KEY_COLUMN = 6
_ALLOWED_MEALS = set(REGULAR_MEALS) | set(RAMADAN_MEALS)
_TEAL = "1D9E75"
_DEEP_TEAL = "085041"
_PALE_TEAL = "E1F5EE"
_INPUT_FILL = "FFF4CC"
_HOLIDAY_FILL = "E5E7EB"
_WHITE = "FFFFFF"
_BORDER = "D1D5DB"


@dataclass(frozen=True)
class BulkWorkbookImport:
    entries: tuple[BulkDailyEntry, ...]
    blank_absence_count: int

    @property
    def first_date(self) -> str:
        return min(entry.date for entry in self.entries)

    @property
    def last_date(self) -> str:
        return max(entry.date for entry in self.entries)


def _date_range(start: date, end: date) -> Iterator[date]:
    current = start
    while current <= end:
        yield current
        current += timedelta(days=1)


def create_bulk_workbook(path: Path, start: date, end: date) -> int:
    """Create a formatted template and return its editable meal-row count."""
    if end < start:
        raise ValueError("تاريخ النهاية يجب أن يكون بعد تاريخ البداية.")

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET_TITLE
    _configure_sheet(sheet)
    _write_headings(sheet)

    holidays = {item.date: item.label for item in get_all_holidays()}
    settings = get_school_settings()
    overrides = get_ramadan_overrides()
    row = _DATA_ROW
    meal_rows = 0
    for day in _date_range(start, end):
        date_text = day.isoformat()
        if date_text in holidays:
            _write_holiday_row(sheet, row, day, holidays[date_text])
            row += 1
            continue
        contacts = {item.meal_type: item for item in get_day_contacts(date_text)}
        absences = {item.meal_type: item for item in get_day_absences(date_text)}
        for meal_type in meals_for_date(date_text, settings, overrides):
            _write_meal_row(
                sheet, row, day, meal_type,
                contacts.get(meal_type), absences.get(meal_type),
            )
            row += 1
            meal_rows += 1

    _finish_sheet(sheet, row - 1)
    _add_metadata(workbook, start, end)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    workbook.close()
    return meal_rows


def _configure_sheet(sheet: Worksheet) -> None:
    sheet.sheet_view.rightToLeft = True
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = f"A{_DATA_ROW}"
    widths = {"A": 15, "B": 13, "C": 20, "D": 16, "E": 16, "F": 2}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    sheet.column_dimensions["F"].hidden = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True


def _write_headings(sheet: Worksheet) -> None:
    for row in range(1, 4):
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
    sheet["A1"] = "إدخال الحضور والغياب لعدة أيام"
    sheet["A2"] = "الإعدادي - منحة كاملة"
    sheet["A3"] = "ملاحظة: إذا تركت خانة الغياب فارغة فسيعتبرها البرنامج 0، ويمكن تغييرها لاحقاً."

    sheet["A1"].font = Font(name="Arial", size=16, bold=True, color=_WHITE)
    sheet["A1"].fill = PatternFill("solid", fgColor=_DEEP_TEAL)
    sheet["A2"].font = Font(name="Arial", size=12, bold=True, color=_DEEP_TEAL)
    sheet["A2"].fill = PatternFill("solid", fgColor=_PALE_TEAL)
    sheet["A3"].font = Font(name="Arial", size=11, bold=True, color="9A3412")
    sheet["A3"].fill = PatternFill("solid", fgColor="FFF7ED")
    for row in range(1, 4):
        sheet.cell(row, 1).alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 30
    sheet.row_dimensions[2].height = 24
    sheet.row_dimensions[3].height = 28

    for column, heading in enumerate(_HEADERS, start=1):
        cell = sheet.cell(_HEADER_ROW, column, heading)
        cell.font = Font(name="Arial", size=11, bold=True, color=_WHITE)
        cell.fill = PatternFill("solid", fgColor=_TEAL)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[_HEADER_ROW].height = 26


def _write_holiday_row(
    sheet: Worksheet, row: int, day: date, label: str,
) -> None:
    values = (day, ARABIC_DAY_NAMES[day.weekday()], f"عطلة: {label}", None, None)
    for column, value in enumerate(values, start=1):
        cell = sheet.cell(row, column, value)
        cell.fill = PatternFill("solid", fgColor=_HOLIDAY_FILL)
        _style_data_cell(cell)
    sheet.cell(row, 1).number_format = "yyyy-mm-dd"
    sheet.cell(row, _MEAL_KEY_COLUMN, "")


def _write_meal_row(
    sheet: Worksheet,
    row: int,
    day: date,
    meal_type: str,
    contact: DailyContact | None,
    absence: DailyAbsence | None,
) -> None:
    attendance = contact.collegial_granted if contact is not None else None
    absence_count = absence.collegial_granted if absence is not None else None
    values = (
        day, ARABIC_DAY_NAMES[day.weekday()], MEAL_LABELS.get(meal_type, meal_type),
        attendance, absence_count,
    )
    for column, value in enumerate(values, start=1):
        cell = sheet.cell(row, column, value)
        _style_data_cell(cell)
        if column in (4, 5):
            cell.fill = PatternFill("solid", fgColor=_INPUT_FILL)
            cell.number_format = "0"
    sheet.cell(row, 1).number_format = "yyyy-mm-dd"
    sheet.cell(row, _MEAL_KEY_COLUMN, meal_type)


def _style_data_cell(cell: Cell) -> None:
    thin = Side(style="thin", color=_BORDER)
    cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    cell.font = Font(name="Arial", size=11, color=_DEEP_TEAL)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.parent.row_dimensions[cell.row].height = 23


def _finish_sheet(sheet: Worksheet, last_row: int) -> None:
    if last_row >= _DATA_ROW:
        validation = DataValidation(
            type="whole", operator="greaterThanOrEqual", formula1="0",
            allow_blank=True,
        )
        validation.error = "أدخل عدداً صحيحاً يساوي 0 أو أكثر."
        validation.errorTitle = "رقم غير صالح"
        sheet.add_data_validation(validation)
        validation.add(f"D{_DATA_ROW}:E{last_row}")
        sheet.auto_filter.ref = f"A{_HEADER_ROW}:E{last_row}"
        sheet.print_area = f"A1:E{last_row}"


def _add_metadata(workbook: Workbook, start: date, end: date) -> None:
    metadata = workbook.create_sheet(_META_SHEET)
    metadata["A1"] = _SIGNATURE
    metadata["A2"] = "collegial_full"
    metadata["A3"] = start.isoformat()
    metadata["A4"] = end.isoformat()
    metadata.sheet_state = "hidden"


def load_bulk_workbook(path: Path) -> BulkWorkbookImport:
    """Read an exported template; blank absence cells become zero."""
    workbook = load_workbook(path, data_only=True, read_only=False)
    try:
        if _META_SHEET not in workbook.sheetnames:
            raise ValueError("هذا الملف ليس نموذج الإدخال اليومي الصادر من البرنامج.")
        metadata = workbook[_META_SHEET]
        if metadata["A1"].value != _SIGNATURE:
            raise ValueError("إصدار نموذج Excel غير معروف.")
        if SHEET_TITLE not in workbook.sheetnames:
            raise ValueError("تعذر العثور على ورقة الإدخال اليومي.")
        return _read_entries(workbook[SHEET_TITLE])
    finally:
        workbook.close()


def _read_entries(sheet: Worksheet) -> BulkWorkbookImport:
    entries: list[BulkDailyEntry] = []
    blank_absences = 0
    for row in range(_DATA_ROW, sheet.max_row + 1):
        meal_type = str(sheet.cell(row, _MEAL_KEY_COLUMN).value or "").strip()
        attendance_raw = sheet.cell(row, 4).value
        absence_raw = sheet.cell(row, 5).value
        if not meal_type:
            continue
        if meal_type not in _ALLOWED_MEALS:
            raise ValueError(f"رمز الوجبة غير صالح في السطر {row}.")
        if _is_blank(attendance_raw):
            if not _is_blank(absence_raw):
                raise ValueError(f"أدخل الحضور أولاً في السطر {row}.")
            continue
        attendance = _count(attendance_raw, row, "الحضور")
        if _is_blank(absence_raw):
            absence = 0
            blank_absences += 1
        else:
            absence = _count(absence_raw, row, "الغياب")
        day = _excel_date(sheet.cell(row, 1).value, row)
        entries.append(BulkDailyEntry(day.isoformat(), meal_type, attendance, absence))

    validate_bulk_daily_entries(entries)
    return BulkWorkbookImport(tuple(entries), blank_absences)


def _is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _count(value: Any, row: int, label: str) -> int:
    if isinstance(value, bool):
        raise ValueError(f"قيمة {label} غير صالحة في السطر {row}.")
    if isinstance(value, int):
        number = value
    elif isinstance(value, float) and value.is_integer():
        number = int(value)
    else:
        try:
            number = int(str(value).strip())
        except ValueError as exc:
            raise ValueError(f"قيمة {label} غير صالحة في السطر {row}.") from exc
    if number < 0:
        raise ValueError(f"قيمة {label} سالبة في السطر {row}.")
    return number


def _excel_date(value: Any, row: int) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    try:
        return date.fromisoformat(str(value).strip())
    except ValueError as exc:
        raise ValueError(f"التاريخ غير صالح في السطر {row}.") from exc
