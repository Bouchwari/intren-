"""تقييم التلاميذ — recorded opinions about the meals served.

The rules this feature has to hold: a dish with no ratings has NO rating (not
zero), every average carries the number of opinions behind it, and a single
five-star rating never makes a dish "the most popular".
"""
import datetime
import sys
import tempfile
import unittest
from pathlib import Path

from PySide6.QtCore import QDate
from PySide6.QtWidgets import QApplication, QMessageBox

ROOT_DIR = Path(__file__).resolve().parents[1]
SRC_DIR = ROOT_DIR / "src"

sys.path.insert(0, str(SRC_DIR))
sys.path.insert(0, str(ROOT_DIR))

from config.settings import MEAL_ASHA, MEAL_GHADA, MEAL_IFTAR, MEAL_SHOUR
from core.feedback import (
    COUNT_FIELDS, MIN_RATINGS_FOR_RANKING, MIN_RESPONSES_FOR_PATTERN,
    RATING_LEVELS, average_by_meal_type, biggest_gaps, favourite_per_group,
    is_valid_rating, least_popular, most_popular, normalize_dish,
    overall_average, record_average, response_counts, response_total,
    summarize_by_attribute, summarize_by_dish, ungrouped_responses,
)
from core.student_stats import profile_by_cycle
from core.models import (
    MealEntry, MealFeedback, SchoolSettings, Student, WeekFeedback,
)
from data import database
from ui.feedback_export import write_feedback_report_pdf
from ui import feedback_collect as fbc
from ui import feedback_screen as fbs


def _rating(dish: str, rating: int, meal: str = MEAL_GHADA,
            date: str = "2026-08-01") -> MealFeedback:
    """A LEGACY row — one single rating, as saved before the counts existed."""
    return MealFeedback(date=date, meal_type=meal, dish=dish, rating=rating)


def _counted(dish: str, counts: dict, meal: str = MEAL_GHADA,
             date: str = "2026-08-01") -> MealFeedback:
    """A row in the current shape: how many pupils gave each level."""
    return MealFeedback(
        date=date, meal_type=meal, dish=dish,
        **{COUNT_FIELDS[level]: value for level, value in counts.items()})


class RatingValidityTests(unittest.TestCase):
    def test_only_one_to_five_is_a_rating(self) -> None:
        self.assertTrue(all(is_valid_rating(v) for v in (1, 2, 3, 4, 5)))
        self.assertFalse(is_valid_rating(0))
        self.assertFalse(is_valid_rating(6))
        self.assertFalse(is_valid_rating(-1))

    def test_invalid_ratings_are_ignored_not_counted_as_zero(self) -> None:
        """A row with rating 0 is not an opinion — it must not drag an
        average down as though someone gave the worst possible score."""
        records = [_rating("كسكس", 4), _rating("كسكس", 0)]
        summaries = summarize_by_dish(records)

        self.assertEqual(len(summaries), 1)
        self.assertEqual(summaries[0].count, 1)
        self.assertEqual(summaries[0].average, 4)


class AggregationTests(unittest.TestCase):
    def test_no_ratings_gives_none_not_zero(self) -> None:
        """"nobody has rated this" and "everyone rated it zero" are different
        claims, and only one of them is possible."""
        self.assertIsNone(overall_average([]))

    def test_the_overall_average_is_the_mean_of_valid_ratings(self) -> None:
        records = [_rating("أ", 5), _rating("ب", 3), _rating("ج", 4)]
        self.assertEqual(overall_average(records), 4)

    def test_dishes_are_grouped_on_the_normalised_name(self) -> None:
        records = [_rating("كسكس بالخضر", 5), _rating("  كسكس   بالخضر ", 3)]
        summaries = summarize_by_dish(records)

        self.assertEqual(len(summaries), 1)
        self.assertEqual(summaries[0].count, 2)
        self.assertEqual(summaries[0].average, 4)

    def test_summaries_carry_their_own_count(self) -> None:
        """An average without its count is not a claim anyone can weigh."""
        records = [_rating("كسكس", 5)] * 4 + [_rating("عدس", 5)]
        summaries = {s.dish: s for s in summarize_by_dish(records)}

        self.assertEqual(summaries["كسكس"].count, 4)
        self.assertEqual(summaries["عدس"].count, 1)
        self.assertEqual(summaries["كسكس"].average, summaries["عدس"].average)

    def test_a_rating_with_no_dish_name_is_skipped(self) -> None:
        self.assertEqual(summarize_by_dish([_rating("   ", 5)]), [])

    def test_average_by_meal_type_separates_the_meals(self) -> None:
        records = [
            _rating("أ", 5, meal=MEAL_GHADA), _rating("ب", 3, meal=MEAL_GHADA),
            _rating("ج", 2, meal=MEAL_ASHA),
        ]
        by_meal = average_by_meal_type(records)

        self.assertEqual(by_meal[MEAL_GHADA].average, 4)
        self.assertEqual(by_meal[MEAL_GHADA].count, 2)
        self.assertEqual(by_meal[MEAL_ASHA].average, 2)


class RankingTests(unittest.TestCase):
    def test_one_rating_never_makes_a_dish_the_most_popular(self) -> None:
        records = [_rating("نادر", 5)] + [_rating("كسكس", 4)] * MIN_RATINGS_FOR_RANKING
        top = most_popular(records)

        self.assertEqual([s.dish for s in top], ["كسكس"])

    def test_nothing_ranks_until_a_dish_has_enough_ratings(self) -> None:
        records = [_rating("كسكس", 5)] * (MIN_RATINGS_FOR_RANKING - 1)
        self.assertEqual(most_popular(records), [])
        self.assertEqual(least_popular(records), [])

    def test_least_popular_is_the_worst_rankable_dish(self) -> None:
        records = (
            [_rating("كسكس", 5)] * MIN_RATINGS_FOR_RANKING
            + [_rating("عدس", 2)] * MIN_RATINGS_FOR_RANKING
        )
        self.assertEqual([s.dish for s in most_popular(records)][0], "كسكس")
        self.assertEqual([s.dish for s in least_popular(records)][0], "عدس")


class FeedbackRepositoryTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.app = QApplication.instance() or QApplication([])

    def setUp(self) -> None:
        self._tmpdir = tempfile.TemporaryDirectory()
        self._original_db_path = database.DB_PATH
        database.DB_PATH = Path(self._tmpdir.name) / "test_matama.db"
        database.init_database()

    def tearDown(self) -> None:
        database.DB_PATH = self._original_db_path
        self._tmpdir.cleanup()

    def test_round_trip_and_delete(self) -> None:
        record_id = database.save_feedback(MealFeedback(
            date="2026-08-01", meal_type=MEAL_GHADA, dish="كسكس", rating=5,
            note="أعجبهم", recorded_by="الحارس العام"))

        everyone = database.get_all_feedback()
        self.assertEqual(len(everyone), 1)
        self.assertEqual(everyone[0].dish, "كسكس")
        self.assertEqual(everyone[0].note, "أعجبهم")

        database.delete_feedback(record_id)
        self.assertEqual(database.get_all_feedback(), [])

    def test_month_filter_returns_only_that_month(self) -> None:
        database.save_feedback(_rating("كسكس", 5, date="2026-08-10"))
        database.save_feedback(_rating("عدس", 3, date="2026-07-10"))

        august = database.get_feedback_for_month("2026-08")
        self.assertEqual([r.dish for r in august], ["كسكس"])


class WeekFeedbackScreenTests(unittest.TestCase):
    """The screen rates a WEEK'S MENU: one row per distinct dish, filled in
    together. The user asked for this after finding day-by-day too much work."""

    @classmethod
    def setUpClass(cls) -> None:
        cls.app = QApplication.instance() or QApplication([])

    def setUp(self) -> None:
        self._tmpdir = tempfile.TemporaryDirectory()
        self._original_db_path = database.DB_PATH
        database.DB_PATH = Path(self._tmpdir.name) / "test_matama.db"
        database.init_database()
        database.save_school_settings(SchoolSettings(
            school_name="مؤسسة", school_year="2025-2026", director="المدير",
            ramadan_start="2026-03-01", ramadan_end="2026-03-30"))
        self._info = QMessageBox.information
        QMessageBox.information = staticmethod(lambda *a, **k: None)

        program_id = database.create_program("أسبوع", "2025-2026", False)
        database.save_program_entries(program_id, [
            MealEntry(program_id=program_id, day_of_week=day,
                      meal_type=MEAL_GHADA, menu_text=text)
            for day, text in ((2, "كسكس بالخضر"), (3, "عدس وخبز"),
                              (4, "كسكس بالخضر"))     # served twice on purpose
        ])

    def tearDown(self) -> None:
        QMessageBox.information = self._info
        database.DB_PATH = self._original_db_path
        self._tmpdir.cleanup()

    def _screen(self, day=QDate(2026, 5, 13)):
        screen = fbs.FeedbackScreen()
        screen.show()
        screen._week_input.setDate(day)
        self.app.processEvents()
        return screen

    def test_any_day_of_the_week_resolves_to_the_same_week(self) -> None:
        for day in (QDate(2026, 5, 11), QDate(2026, 5, 14), QDate(2026, 5, 17)):
            screen = self._screen(day)
            self.assertEqual(screen._selected_week(), "2026-05-11")
            screen.close()

    def test_a_dish_served_twice_appears_once(self) -> None:
        """The whole point of rating the menu rather than the days."""
        screen = self._screen()
        dishes = [screen._menu_table.item(row, 0).text()
                  for row in range(screen._menu_table.rowCount())]

        self.assertEqual(sorted(dishes), ["عدس وخبز", "كسكس بالخضر"])
        screen.close()

    def test_saving_writes_one_row_per_rated_dish(self) -> None:
        screen = self._screen()
        spins = screen._menu_spins["كسكس بالخضر"]
        spins[5].setValue(60)
        spins[4].setValue(40)
        screen._on_save_week()

        saved = database.get_week_feedback("2026-05-11")
        self.assertEqual([r.dish for r in saved], ["كسكس بالخضر"])
        self.assertEqual(response_total(saved[0]), 100)
        screen.close()

    def test_a_dish_left_at_zero_is_not_saved(self) -> None:
        """A dish nobody was asked about must not become a dish nobody liked."""
        screen = self._screen()
        screen._menu_spins["كسكس بالخضر"][5].setValue(10)
        screen._on_save_week()

        self.assertEqual(len(database.get_week_feedback("2026-05-11")), 1)
        screen.close()

    def test_saving_the_same_week_twice_updates_it(self) -> None:
        screen = self._screen()
        screen._menu_spins["عدس وخبز"][1].setValue(50)
        screen._on_save_week()
        screen._menu_spins["عدس وخبز"][1].setValue(30)
        screen._on_save_week()

        saved = database.get_week_feedback("2026-05-11")
        self.assertEqual(len(saved), 1)
        self.assertEqual(response_total(saved[0]), 30)
        screen.close()

    def test_saved_counts_come_back_when_the_week_is_reopened(self) -> None:
        screen = self._screen()
        screen._menu_spins["كسكس بالخضر"][5].setValue(70)
        screen._on_save_week()
        screen.close()

        reopened = self._screen()
        self.assertEqual(reopened._menu_spins["كسكس بالخضر"][5].value(), 70)
        reopened.close()

    def test_a_dish_rated_before_but_no_longer_on_the_menu_keeps_its_row(self) -> None:
        """Otherwise its numbers would silently vanish from the screen that
        owns them."""
        database.save_week_feedback(WeekFeedback(
            week_start="2026-05-11", dish="طبق قديم", count_good=20))

        screen = self._screen()
        dishes = [screen._menu_table.item(row, 0).text()
                  for row in range(screen._menu_table.rowCount())]
        self.assertIn("طبق قديم", dishes)
        screen.close()

    def test_the_old_per_day_records_are_kept_and_still_counted(self) -> None:
        """The user asked for them to be kept and shown separately."""
        database.save_feedback(_counted("كسكس بالخضر", {5: 10}, date="2026-04-06"))
        screen = self._screen()

        self.assertEqual(len(screen._records), 1)
        self.assertEqual(
            sum(response_total(r) for r in screen._all_opinions()), 10)
        screen.close()

    def test_a_ramadan_week_lists_its_own_dishes(self) -> None:
        ramadan_id = database.create_program("رمضان", "2025-2026", True)
        database.save_program_entries(ramadan_id, [
            MealEntry(program_id=ramadan_id, day_of_week=2,
                      meal_type=MEAL_IFTAR, menu_text="حريرة وتمر"),
        ])
        screen = self._screen(QDate(2026, 3, 4))

        dishes = [screen._menu_table.item(row, 0).text()
                  for row in range(screen._menu_table.rowCount())]
        self.assertIn("حريرة وتمر", dishes)
        self.assertNotIn("كسكس بالخضر", dishes)
        screen.close()


class ResponseCountTests(unittest.TestCase):
    """A meal is eaten by a whole school, so what gets recorded is how many
    pupils said what — one opinion per meal was never a measurement."""

    def test_counts_are_read_back_per_level(self) -> None:
        record = _counted("كسكس", {5: 40, 4: 55, 3: 20, 2: 8, 1: 2})
        self.assertEqual(response_counts(record), {5: 40, 4: 55, 3: 20, 2: 8, 1: 2})
        self.assertEqual(response_total(record), 125)
        self.assertAlmostEqual(record_average(record), 3.984)

    def test_a_legacy_single_rating_reads_as_exactly_one_response(self) -> None:
        """Rows saved before the counts existed still count for something,
        without being inflated into a crowd."""
        record = _rating("كسكس", 4)
        self.assertEqual(response_total(record), 1)
        self.assertEqual(record_average(record), 4)

    def test_counts_win_over_a_legacy_rating_on_the_same_row(self) -> None:
        record = _counted("كسكس", {5: 10})
        record.rating = 1
        self.assertEqual(response_total(record), 10)
        self.assertEqual(record_average(record), 5)

    def test_a_row_with_nobody_in_it_has_no_average(self) -> None:
        record = MealFeedback(date="2026-08-01", meal_type=MEAL_GHADA, dish="كسكس")
        self.assertEqual(response_total(record), 0)
        self.assertIsNone(record_average(record))

    def test_negative_counts_are_treated_as_none(self) -> None:
        record = _counted("كسكس", {5: 10})
        record.count_bad = -5
        self.assertEqual(response_total(record), 10)

    def test_the_dish_average_is_weighted_by_how_many_gave_each_level(self) -> None:
        """Not the mean of the levels — the mean of the OPINIONS."""
        summary = summarize_by_dish([_counted("كسكس", {5: 90, 1: 10})])[0]
        self.assertEqual(summary.count, 100)
        self.assertAlmostEqual(summary.average, 4.6)

    def test_ranking_uses_responses_not_rows(self) -> None:
        """One row carrying 100 responses is plenty to rank on; three rows
        carrying one opinion each is the case the threshold guards against."""
        rankable = most_popular([_counted("كسكس", {5: 100})])
        self.assertEqual([s.dish for s in rankable], ["كسكس"])

        self.assertEqual(most_popular([_counted("عدس", {5: 1})]), [])

    def test_counted_and_legacy_rows_combine(self) -> None:
        records = [_counted("كسكس", {4: 9}), _rating("كسكس", 5)]
        summary = summarize_by_dish(records)[0]
        self.assertEqual(summary.count, 10)
        self.assertAlmostEqual(summary.average, 4.1)

    def test_the_overall_average_weights_every_response(self) -> None:
        records = [_counted("أ", {5: 90}), _counted("ب", {1: 10})]
        self.assertAlmostEqual(overall_average(records), 4.6)


class FeedbackCollectionTests(unittest.TestCase):
    """Getting opinions off the refectory floor: a printed tally sheet and an
    Excel round-trip, both listing the WEEK'S DISHES. The QR/phone route was
    deliberately NOT built — the pupils have phones but no WiFi reaching the
    office PC."""

    @classmethod
    def setUpClass(cls) -> None:
        cls.app = QApplication.instance() or QApplication([])

    def setUp(self) -> None:
        self._tmpdir = tempfile.TemporaryDirectory()
        self._original_db_path = database.DB_PATH
        database.DB_PATH = Path(self._tmpdir.name) / "test_matama.db"
        database.init_database()
        self._settings = SchoolSettings(
            school_name="مؤسسة", school_year="2025-2026", director="المدير")
        database.save_school_settings(self._settings)

        program_id = database.create_program("أسبوع", "2025-2026", False)
        database.save_program_entries(program_id, [
            MealEntry(program_id=program_id, day_of_week=day,
                      meal_type=MEAL_GHADA, menu_text=text)
            for day, text in ((2, "كسكس بالخضر"), (3, "عدس وخبز"),
                              (4, "كسكس بالخضر"))
        ])
        self._week = "2026-05-11"

    def tearDown(self) -> None:
        database.DB_PATH = self._original_db_path
        self._tmpdir.cleanup()

    def _dishes(self):
        from ui.feedback_collect import dishes_for_week
        return dishes_for_week(self._week, database.get_school_settings())

    def test_the_week_lists_each_dish_once(self) -> None:
        self.assertEqual(sorted(self._dishes()), ["عدس وخبز", "كسكس بالخضر"])

    def test_the_week_label_reads_as_a_date_range(self) -> None:
        from ui.feedback_collect import week_label
        self.assertEqual(week_label("2026-05-13"), "أسبوع 11 - 17 مايو 2026")

    def test_the_tally_sheet_is_a_real_pdf(self) -> None:
        from ui.feedback_collect import write_tally_sheet_pdf

        path = Path(self._tmpdir.name) / "tally.pdf"
        write_tally_sheet_pdf(path, self._settings, self._dishes(), self._week)

        self.assertEqual(path.read_bytes()[:4], b"%PDF")
        self.assertGreater(path.stat().st_size, 2000)

    def _workbook(self) -> Path:
        from ui.feedback_collect import write_feedback_workbook
        path = Path(self._tmpdir.name) / "feedback.xlsx"
        write_feedback_workbook(path, self._dishes(), self._week)
        return path

    def test_the_workbook_carries_the_week_and_empty_counts(self) -> None:
        """Blank, not zero: an empty cell means "not collected"."""
        import openpyxl

        sheet = openpyxl.load_workbook(self._workbook()).active
        self.assertEqual(str(sheet["B2"].value), self._week)
        self.assertEqual(sheet.cell(row=3, column=2).value, "ممتاز")
        for column in range(2, 7):
            self.assertIsNone(sheet.cell(row=4, column=column).value)

    def _fill(self, path: Path, row: int, counts) -> None:
        import openpyxl
        workbook = openpyxl.load_workbook(path)
        sheet = workbook.active
        for offset, value in enumerate(counts):
            sheet.cell(row=row, column=2 + offset, value=value)
        workbook.save(path)

    def test_a_filled_dish_is_imported_against_its_week(self) -> None:
        from ui.feedback_collect import import_feedback_workbook

        path = self._workbook()
        self._fill(path, 4, [40, 55, 20, 8, 2])

        saved, _skipped, problems, week = import_feedback_workbook(path)
        self.assertEqual((saved, problems, week), (1, [], self._week))
        self.assertEqual(response_total(database.get_week_feedback(self._week)[0]), 125)

    def test_dishes_left_blank_are_skipped_not_stored_as_zero(self) -> None:
        from ui.feedback_collect import import_feedback_workbook

        path = self._workbook()
        self._fill(path, 4, [10, 0, 0, 0, 0])

        saved, skipped, _problems, _week = import_feedback_workbook(path)
        self.assertEqual((saved, skipped), (1, 1))
        self.assertEqual(len(database.get_week_feedback(self._week)), 1)

    def test_an_unreadable_count_is_reported_never_guessed(self) -> None:
        from ui.feedback_collect import import_feedback_workbook

        path = self._workbook()
        self._fill(path, 4, ["كثير", 0, 0, 0, 0])

        saved, _skipped, problems, _week = import_feedback_workbook(path)
        self.assertEqual(saved, 0)
        self.assertEqual(len(problems), 1)

    def test_a_negative_count_is_refused(self) -> None:
        from ui.feedback_collect import import_feedback_workbook

        path = self._workbook()
        self._fill(path, 4, [-5, 0, 0, 0, 0])

        saved, _skipped, problems, _week = import_feedback_workbook(path)
        self.assertEqual((saved, len(problems)), (0, 1))

    def test_a_file_without_a_week_is_refused(self) -> None:
        """B2 carries the week; without it the counts belong to nothing."""
        import openpyxl
        from ui.feedback_collect import import_feedback_workbook

        path = self._workbook()
        workbook = openpyxl.load_workbook(path)
        workbook.active["B2"] = ""
        workbook.save(path)

        saved, _skipped, problems, week = import_feedback_workbook(path)
        self.assertEqual((saved, week), (0, ""))
        self.assertEqual(len(problems), 1)

    def test_importing_the_same_sheet_twice_does_not_double_the_counts(self) -> None:
        """The most likely user mistake — pressing import again."""
        from ui.feedback_collect import import_feedback_workbook

        path = self._workbook()
        self._fill(path, 4, [40, 55, 20, 8, 2])
        import_feedback_workbook(path)
        import_feedback_workbook(path)

        saved = database.get_week_feedback(self._week)
        self.assertEqual(len(saved), 1)
        self.assertEqual(response_total(saved[0]), 125)

    def test_week_records_are_skipped_by_the_per_meal_breakdown(self) -> None:
        """A week rating covers a dish across the whole week, so it cannot be
        filed under فطور or غداء — it must be skipped, not crash."""
        records = [
            _counted("كسكس", {5: 10}, meal=MEAL_GHADA),
            WeekFeedback(week_start="2026-05-11", dish="عدس", count_bad=20),
        ]
        by_meal = average_by_meal_type(records)

        self.assertEqual(set(by_meal), {MEAL_GHADA})
        self.assertEqual(by_meal[MEAL_GHADA].count, 10)


class ReportContentTests(unittest.TestCase):
    """The report's indicators and charts."""

    @classmethod
    def setUpClass(cls) -> None:
        cls.app = QApplication.instance() or QApplication([])

    def setUp(self) -> None:
        self._tmpdir = tempfile.TemporaryDirectory()
        self._settings = SchoolSettings(
            school_name="مؤسسة", school_year="2025-2026", director="المدير")

    def tearDown(self) -> None:
        self._tmpdir.cleanup()

    def test_positive_and_negative_shares_split_at_the_middle(self) -> None:
        """4-5 counts as positive, 1-2 as negative, 3 as neither."""
        from ui.feedback_export import _positive_negative

        records = [_counted("كسكس", {5: 10, 4: 20, 3: 30, 2: 5, 1: 5})]
        positive, negative, total = _positive_negative(records)

        self.assertEqual((positive, negative, total), (30, 10, 70))

    def test_the_weekly_trend_uses_week_records_only(self) -> None:
        """A legacy per-day row has no week, so it cannot sit on a week axis."""
        from ui.feedback_export import _weekly_averages

        records = [
            WeekFeedback(week_start="2026-05-11", dish="أ", count_excellent=10),
            WeekFeedback(week_start="2026-05-18", dish="ب", count_bad=10),
            _counted("ج", {5: 10}),          # legacy, no week
        ]
        trend = _weekly_averages(records)

        self.assertEqual([label for label, _v in trend], ["05-11", "05-18"])
        self.assertEqual([round(v, 2) for _l, v in trend], [5.0, 1.0])

    def test_the_report_adds_an_indicator_page_and_a_population_page(self) -> None:
        from ui.feedback_export import write_feedback_report_pdf

        records = [WeekFeedback(week_start="2026-05-11", dish="كسكس",
                                count_excellent=40, count_good=20)]
        students = [Student(full_name="أ", gender="male", cycle="إعدادي",
                            birth_date="2012-01-01")]
        path = Path(self._tmpdir.name) / "report.pdf"
        write_feedback_report_pdf(path, self._settings, records, "فترة",
                                  students=students,
                                  today=datetime.date(2026, 8, 29))

        self.assertEqual(path.read_bytes()[:4], b"%PDF")
        # page 1 (kept as it was) + indicators + population
        self.assertGreaterEqual(path.read_bytes().count(b"/Type /Page"), 3)

    def test_with_no_students_the_population_page_is_omitted(self) -> None:
        """Printing an empty roster page would say nothing."""
        from ui.feedback_export import write_feedback_report_pdf

        records = [WeekFeedback(week_start="2026-05-11", dish="كسكس",
                                count_good=10)]
        with_students = Path(self._tmpdir.name) / "with.pdf"
        without = Path(self._tmpdir.name) / "without.pdf"
        write_feedback_report_pdf(with_students, self._settings, records, "فترة",
                                  students=[Student(full_name="أ")],
                                  today=datetime.date(2026, 8, 29))
        write_feedback_report_pdf(without, self._settings, records, "فترة",
                                  students=[], today=datetime.date(2026, 8, 29))

        self.assertGreater(with_students.read_bytes().count(b"/Type /Page"),
                           without.read_bytes().count(b"/Type /Page"))


# ── Who gave the opinions: cycle × gender patterns ──────────────────────────

def _pdf_pages(path: Path) -> int:
    """How many pages a generated PDF really has."""
    import re
    counts = re.findall(rb"/Count\s+(\d+)", path.read_bytes())
    return int(counts[0]) if counts else 0


def _group(dish: str, cycle: str, gender: str, counts: dict,
           week: str = "2026-05-11") -> WeekFeedback:
    record = WeekFeedback(week_start=week, dish=dish, cycle=cycle,
                          gender=gender)
    for level, field in COUNT_FIELDS.items():
        setattr(record, field, counts.get(level, 0))
    return record


class GroupPatternTests(unittest.TestCase):
    """The report may say "تأهيلي rate this higher than ابتدائي" only when
    both groups really answered enough to be compared."""

    def test_records_with_no_group_are_left_out_of_the_comparison(self) -> None:
        records = [_group("كسكس", "primary", "male", {5: 10}),
                   WeekFeedback(week_start="2026-05-11", dish="كسكس",
                                count_good=40)]
        by_cycle = summarize_by_attribute(records, "cycle")
        self.assertEqual(list(by_cycle), ["primary"])
        self.assertEqual(by_cycle["primary"].count, 10)
        # …but they are still reported, so the reader knows what is missing.
        self.assertEqual(ungrouped_responses(records), 40)

    def test_a_dish_only_one_group_rated_is_not_a_disagreement(self) -> None:
        records = [_group("عدس", "primary", "male", {1: 30})]
        self.assertEqual(biggest_gaps(records, "cycle"), [])

    def test_a_group_below_the_minimum_is_not_compared(self) -> None:
        few = MIN_RESPONSES_FOR_PATTERN - 1
        records = [_group("عدس", "primary", "male", {1: 30}),
                   _group("عدس", "qualifying", "female", {5: few})]
        self.assertEqual(biggest_gaps(records, "cycle"), [])

    def test_gaps_come_back_widest_first_with_both_sides_named(self) -> None:
        records = [
            _group("عدس", "primary", "male", {1: 40}),        # 1.00
            _group("عدس", "qualifying", "female", {5: 40}),   # 5.00
            _group("كسكس", "primary", "male", {4: 40}),       # 4.00
            _group("كسكس", "qualifying", "female", {3: 40}),  # 3.00
        ]
        gaps = biggest_gaps(records, "cycle")
        self.assertEqual([gap.dish for gap in gaps], ["عدس", "كسكس"])
        self.assertEqual(gaps[0].high_group, "qualifying")
        self.assertEqual(gaps[0].low_group, "primary")
        self.assertAlmostEqual(gaps[0].gap, 4.0)
        # The count travels with the average — an average alone is not a claim
        # anyone can weigh.
        self.assertEqual(gaps[0].high_count, 40)

    def test_a_group_favourite_needs_enough_opinions(self) -> None:
        records = [_group("عدس", "primary", "male", {3: 40}),
                   _group("كسكس", "primary", "male", {5: 2})]   # too few
        favourites = favourite_per_group(records, "cycle")
        self.assertEqual(favourites["primary"].dish, "عدس")

    def test_gender_uses_the_same_machinery(self) -> None:
        records = [_group("سمك", "primary", "female", {5: 20}),
                   _group("سمك", "primary", "male", {2: 20})]
        gaps = biggest_gaps(records, "gender")
        self.assertEqual(gaps[0].high_group, "female")
        self.assertEqual(gaps[0].low_group, "male")


class GroupProfileTests(unittest.TestCase):
    """The age behind a group comes from the ROSTER, never from the ratings."""

    def _pupil(self, cycle: str, gender: str, birth: str) -> Student:
        return Student(full_name="تلميذ", cycle=cycle, gender=gender,
                       birth_date=birth)

    def test_cycle_profiles_use_the_apps_own_classifier(self) -> None:
        today = datetime.date(2026, 8, 29)
        profiles = profile_by_cycle([
            self._pupil("ابتدائي", "male", "2016-01-01"),
            self._pupil("ابتدائي", "female", "2015-01-01"),
            self._pupil("تأهيلي", "male", "2009-01-01"),
        ], today)
        self.assertEqual(profiles["primary"].count, 2)
        self.assertEqual(profiles["primary"].age_range, (10, 11))
        self.assertAlmostEqual(profiles["qualifying"].average_age, 17.0)

    def test_a_missing_birth_date_is_counted_not_invented(self) -> None:
        profiles = profile_by_cycle(
            [self._pupil("ابتدائي", "male", "")], datetime.date(2026, 8, 29))
        self.assertEqual(profiles["primary"].count, 1)
        self.assertEqual(profiles["primary"].unknown_age, 1)
        self.assertIsNone(profiles["primary"].average_age)


class GroupStorageTests(unittest.TestCase):
    """One row per (week, dish, cycle, gender) — groups must not overwrite
    each other, and re-saving one must not create a second."""

    @classmethod
    def setUpClass(cls) -> None:
        cls.app = QApplication.instance() or QApplication([])

    def setUp(self) -> None:
        self._tmpdir = tempfile.TemporaryDirectory()
        self._original_db_path = database.DB_PATH
        database.DB_PATH = Path(self._tmpdir.name) / "test_matama.db"
        database.init_database()

    def tearDown(self) -> None:
        database.DB_PATH = self._original_db_path
        self._tmpdir.cleanup()

    def test_two_groups_of_the_same_dish_are_two_rows(self) -> None:
        database.save_week_feedback(_group("كسكس", "primary", "male", {5: 10}))
        database.save_week_feedback(
            _group("كسكس", "qualifying", "female", {1: 8}))
        rows = database.get_week_feedback("2026-05-11")
        self.assertEqual(len(rows), 2)
        self.assertEqual({(r.cycle, r.gender) for r in rows},
                         {("primary", "male"), ("qualifying", "female")})

    def test_saving_the_same_group_again_updates_it(self) -> None:
        database.save_week_feedback(_group("كسكس", "primary", "male", {5: 10}))
        database.save_week_feedback(_group("كسكس", "primary", "male", {5: 12}))
        rows = database.get_week_feedback("2026-05-11")
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0].count_excellent, 12)

    def test_an_ungrouped_row_lives_beside_the_grouped_ones(self) -> None:
        database.save_week_feedback(_group("كسكس", "", "", {4: 5}))
        database.save_week_feedback(_group("كسكس", "primary", "male", {5: 10}))
        self.assertEqual(len(database.get_week_feedback("2026-05-11")), 2)

    def test_a_database_from_before_the_grouping_is_rebuilt_not_broken(self) -> None:
        """The old table's UNIQUE(week_start, dish) cannot be widened by ALTER
        TABLE, so init_database rebuilds it — without losing a single row."""
        import sqlite3
        with sqlite3.connect(database.DB_PATH) as conn:
            conn.executescript("""
                DROP TABLE week_feedback;
                CREATE TABLE week_feedback (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    week_start TEXT NOT NULL, dish TEXT NOT NULL,
                    count_excellent INTEGER NOT NULL DEFAULT 0,
                    count_good INTEGER NOT NULL DEFAULT 0,
                    count_average INTEGER NOT NULL DEFAULT 0,
                    count_poor INTEGER NOT NULL DEFAULT 0,
                    count_bad INTEGER NOT NULL DEFAULT 0,
                    note TEXT NOT NULL DEFAULT '',
                    recorded_by TEXT NOT NULL DEFAULT '',
                    created_at TEXT NOT NULL DEFAULT (datetime('now')),
                    UNIQUE(week_start, dish));
                INSERT INTO week_feedback (week_start, dish, count_good)
                VALUES ('2026-05-11', 'قديم', 9);
            """)

        database.init_database()
        database.save_week_feedback(_group("قديم", "primary", "male", {5: 4}))

        rows = database.get_week_feedback("2026-05-11")
        self.assertEqual(len(rows), 2)
        old = next(r for r in rows if not r.cycle)
        self.assertEqual(old.count_good, 9)          # the old row survived
        self.assertEqual(old.gender, "")             # and is غير محدد


class GroupCollectionTests(unittest.TestCase):
    """The paper and Excel routes have to say WHICH group each sheet is for,
    and put the numbers back where they came from."""

    @classmethod
    def setUpClass(cls) -> None:
        cls.app = QApplication.instance() or QApplication([])

    def setUp(self) -> None:
        self._tmpdir = tempfile.TemporaryDirectory()
        self._original_db_path = database.DB_PATH
        database.DB_PATH = Path(self._tmpdir.name) / "test_matama.db"
        database.init_database()
        self._settings = SchoolSettings(
            school_name="مؤسسة", school_year="2025-2026", director="المدير")
        database.save_school_settings(self._settings)
        self._week = "2026-05-11"
        self._dishes = ["كسكس بالخضر", "عدس وخبز"]

    def tearDown(self) -> None:
        database.DB_PATH = self._original_db_path
        self._tmpdir.cleanup()

    def test_no_group_chosen_means_collect_every_group(self) -> None:
        self.assertEqual(len(fbc.groups_to_collect("", "")), 6)
        self.assertEqual(fbc.groups_to_collect("primary", "male"),
                         [("primary", "male")])

    def test_a_group_name_this_app_did_not_write_is_refused(self) -> None:
        self.assertEqual(fbc.parse_group_label("ابتدائي / ذكور"),
                         ("primary", "male"))
        self.assertEqual(fbc.parse_group_label("غير محدد"), ("", ""))
        # Not "assume ابتدائي" — a guessed group is a fabricated opinion.
        self.assertIsNone(fbc.parse_group_label("الصغار"))

    def test_the_workbook_has_one_sheet_per_group(self) -> None:
        from openpyxl import load_workbook
        path = Path(self._tmpdir.name) / "book.xlsx"
        fbc.write_feedback_workbook(path, self._dishes, self._week,
                                    fbc.groups_to_collect("", ""))
        sheets = load_workbook(path).sheetnames
        self.assertEqual(len(sheets), 6)
        self.assertIn("ابتدائي-ذكور", sheets)

    def test_imported_counts_land_in_their_own_group(self) -> None:
        from openpyxl import load_workbook
        path = Path(self._tmpdir.name) / "book.xlsx"
        fbc.write_feedback_workbook(path, self._dishes, self._week,
                                    fbc.groups_to_collect("", ""))
        book = load_workbook(path)
        for sheet_name, counts in (("ابتدائي-ذكور", [9, 1, 0, 0, 0]),
                                   ("تأهيلي-إناث", [0, 0, 0, 2, 7])):
            for column, value in enumerate(counts, start=2):
                book[sheet_name].cell(row=4, column=column, value=value)
        book.save(path)

        saved, _skipped, problems, week = fbc.import_feedback_workbook(path)
        self.assertEqual((saved, problems, week), (2, [], self._week))
        rows = {(r.cycle, r.gender): r
                for r in database.get_week_feedback(self._week)}
        self.assertEqual(rows[("primary", "male")].count_excellent, 9)
        self.assertEqual(rows[("qualifying", "female")].count_bad, 7)

    def test_a_sheet_whose_group_was_edited_is_reported_not_stored(self) -> None:
        from openpyxl import load_workbook
        path = Path(self._tmpdir.name) / "book.xlsx"
        fbc.write_feedback_workbook(path, self._dishes, self._week,
                                    [("primary", "male")])
        book = load_workbook(path)
        sheet = book.worksheets[0]
        sheet["D2"] = "الصغار"
        sheet.cell(row=4, column=2, value=5)
        book.save(path)

        saved, _skipped, problems, _week = fbc.import_feedback_workbook(path)
        self.assertEqual(saved, 0)
        self.assertTrue(problems)
        self.assertEqual(database.get_week_feedback(self._week), [])

    def test_the_tally_sheet_prints_a_page_set_per_group(self) -> None:
        one = Path(self._tmpdir.name) / "one.pdf"
        every = Path(self._tmpdir.name) / "every.pdf"
        fbc.write_tally_sheet_pdf(one, self._settings, self._dishes,
                                  self._week, [("primary", "male")])
        fbc.write_tally_sheet_pdf(every, self._settings, self._dishes,
                                  self._week, fbc.groups_to_collect("", ""))
        # Counted from the PDF itself. File size is no proxy: the fonts are
        # embedded once, so a six-page file is barely larger than a one-page one.
        self.assertEqual(_pdf_pages(one), 1)
        self.assertEqual(_pdf_pages(every), 6)


class GroupScreenTests(unittest.TestCase):
    """Picking a group must change what the table shows AND where a save
    goes — showing one group's numbers under another's name would copy them
    onto the wrong pupils on the next save."""

    @classmethod
    def setUpClass(cls) -> None:
        cls.app = QApplication.instance() or QApplication([])

    def setUp(self) -> None:
        self._tmpdir = tempfile.TemporaryDirectory()
        self._original_db_path = database.DB_PATH
        database.DB_PATH = Path(self._tmpdir.name) / "test_matama.db"
        database.init_database()
        database.save_school_settings(SchoolSettings(
            school_name="مؤسسة", school_year="2025-2026", director="المدير"))
        self._info = QMessageBox.information
        QMessageBox.information = staticmethod(lambda *a, **k: None)
        program_id = database.create_program("أسبوع", "2025-2026", False)
        database.save_program_entries(program_id, [
            MealEntry(program_id=program_id, day_of_week=2,
                      meal_type=MEAL_GHADA, menu_text="كسكس بالخضر")])

    def tearDown(self) -> None:
        QMessageBox.information = self._info
        database.DB_PATH = self._original_db_path
        self._tmpdir.cleanup()

    def _screen(self):
        screen = fbs.FeedbackScreen()
        screen._week_input.setDate(QDate(2026, 5, 13))
        screen._reload()
        return screen

    def _choose(self, screen, cycle: str, gender: str) -> None:
        screen._cycle_combo.setCurrentIndex(
            screen._cycle_combo.findData(cycle))
        screen._gender_combo.setCurrentIndex(
            screen._gender_combo.findData(gender))

    def test_a_save_carries_the_chosen_group(self) -> None:
        screen = self._screen()
        self._choose(screen, "collegial", "female")
        screen._menu_spins["كسكس بالخضر"][5].setValue(7)
        screen._on_save_week()

        rows = database.get_week_feedback("2026-05-11")
        self.assertEqual(len(rows), 1)
        self.assertEqual((rows[0].cycle, rows[0].gender),
                         ("collegial", "female"))
        self.assertEqual(rows[0].count_excellent, 7)

    def test_switching_group_shows_that_group_and_not_the_other(self) -> None:
        database.save_week_feedback(
            _group("كسكس بالخضر", "primary", "male", {5: 11}))
        database.save_week_feedback(
            _group("كسكس بالخضر", "qualifying", "female", {1: 4}))
        screen = self._screen()

        self._choose(screen, "primary", "male")
        self.assertEqual(screen._menu_spins["كسكس بالخضر"][5].value(), 11)
        self.assertEqual(screen._menu_spins["كسكس بالخضر"][1].value(), 0)

        self._choose(screen, "qualifying", "female")
        self.assertEqual(screen._menu_spins["كسكس بالخضر"][5].value(), 0)
        self.assertEqual(screen._menu_spins["كسكس بالخضر"][1].value(), 4)

    def test_an_unchosen_group_starts_empty_not_prefilled(self) -> None:
        """Otherwise the next save would copy one group's answers onto a
        group that never gave them."""
        database.save_week_feedback(
            _group("كسكس بالخضر", "primary", "male", {5: 11}))
        screen = self._screen()
        self._choose(screen, "collegial", "male")
        self.assertEqual(screen._menu_spins["كسكس بالخضر"][5].value(), 0)

    def test_the_screen_lists_which_groups_are_already_collected(self) -> None:
        database.save_week_feedback(
            _group("كسكس بالخضر", "primary", "male", {5: 11}))
        screen = self._screen()
        self.assertIn("ابتدائي / ذكور", screen._groups_label.text())
        self.assertIn("11", screen._groups_label.text())


class PatternsPageTests(unittest.TestCase):
    """The report's "من أبدى الرأي" page. It must appear whether or not the
    collecting has been split by group — saying nothing was split is a
    finding, printing four empty comparison boxes is not."""

    @classmethod
    def setUpClass(cls) -> None:
        cls.app = QApplication.instance() or QApplication([])

    def setUp(self) -> None:
        self._tmpdir = tempfile.TemporaryDirectory()
        self._settings = SchoolSettings(
            school_name="مؤسسة", school_year="2025-2026", director="المدير")
        self._today = datetime.date(2026, 8, 29)
        self._students = [
            Student(full_name="أ", cycle="ابتدائي", gender="male",
                    birth_date="2015-01-01"),
            Student(full_name="ب", cycle="تأهيلي", gender="female",
                    birth_date="2009-01-01"),
        ]

    def tearDown(self) -> None:
        self._tmpdir.cleanup()

    def _write(self, records) -> Path:
        path = Path(self._tmpdir.name) / "report.pdf"
        write_feedback_report_pdf(path, self._settings, records, "أسبوع",
                                  self._students, self._today)
        return path

    def test_the_page_is_written_even_with_no_grouped_opinions(self) -> None:
        path = self._write([_group("كسكس", "", "", {4: 20})])
        self.assertTrue(path.exists())
        self.assertGreaterEqual(_pdf_pages(path), 3)

    def test_grouped_opinions_produce_a_report_without_failing(self) -> None:
        records = [_group("عدس", "primary", "male", {1: 40}),
                   _group("عدس", "qualifying", "female", {5: 40})]
        path = self._write(records)
        self.assertGreater(path.stat().st_size, 0)

    def test_the_report_survives_an_empty_roster(self) -> None:
        """No students imported yet — the age link has nothing behind it and
        must not take the whole report down with it."""
        self._students = []
        path = self._write([_group("عدس", "primary", "male", {1: 40}),
                            _group("عدس", "qualifying", "female", {5: 40})])
        self.assertTrue(path.exists())
