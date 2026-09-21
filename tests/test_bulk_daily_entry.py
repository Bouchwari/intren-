"""Bulk real-number entry through the Excel workflow."""
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook
from PySide6.QtCore import QDate
from PySide6.QtWidgets import QApplication, QDialog, QFileDialog, QLabel, QMessageBox


ROOT_DIR = Path(__file__).resolve().parents[1]
SRC_DIR = ROOT_DIR / "src"
sys.path.insert(0, str(SRC_DIR))
sys.path.insert(0, str(ROOT_DIR))

from config.settings import (
    MEAL_ASHA, MEAL_FTOUR, MEAL_GHADA, MEAL_IFTAR, MEAL_SHOUR,
)
from core.bulk_daily_entry import save_bulk_daily_entries
from core.models import BulkDailyEntry, DailyAbsence, DailyContact, Holiday, SchoolSettings
from data import database
from ui.bulk_daily_entry import BulkDailyEntryDialog
from ui.bulk_daily_excel import SHEET_TITLE, create_bulk_workbook, load_bulk_workbook


class BulkDailyEntryTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.app = QApplication.instance() or QApplication([])

    def setUp(self) -> None:
        self._tmpdir = tempfile.TemporaryDirectory()
        self._original_db_path = database.DB_PATH
        database.DB_PATH = Path(self._tmpdir.name) / "test_matama.db"
        database.init_database()

    def tearDown(self) -> None:
        database.DB_PATH = self._original_db_path
        self._tmpdir.cleanup()

    def _workbook_path(self, name: str = "daily.xlsx") -> Path:
        return Path(self._tmpdir.name) / name

    def test_service_saves_many_meals_atomically(self) -> None:
        save_bulk_daily_entries([
            BulkDailyEntry("2026-06-01", MEAL_FTOUR, 90, 3),
            BulkDailyEntry("2026-06-01", MEAL_GHADA, 95, 2),
            BulkDailyEntry("2026-06-02", MEAL_ASHA, 88, 4),
        ])

        contacts = database.get_day_contacts("2026-06-01")
        absences = database.get_day_absences("2026-06-01")
        self.assertEqual(
            {item.meal_type: item.collegial_granted for item in contacts},
            {MEAL_FTOUR: 90, MEAL_GHADA: 95},
        )
        self.assertEqual(
            {item.meal_type: item.collegial_granted for item in absences},
            {MEAL_FTOUR: 3, MEAL_GHADA: 2},
        )

    def test_bulk_update_preserves_other_official_form_categories(self) -> None:
        database.save_daily_contact(DailyContact(
            date="2026-06-01", meal_type=MEAL_GHADA,
            primary_granted=7, collegial_granted=10,
            collegial_complement=4, qualifying_granted=6,
        ))
        database.save_daily_absence(DailyAbsence(
            date="2026-06-01", meal_type=MEAL_GHADA,
            primary_granted=1, collegial_granted=2,
            qualifying_granted=3,
        ))

        save_bulk_daily_entries([
            BulkDailyEntry("2026-06-01", MEAL_GHADA, 80, 5),
        ])

        contact = database.get_day_contacts("2026-06-01")[0]
        absence = database.get_day_absences("2026-06-01")[0]
        self.assertEqual(contact.collegial_granted, 80)
        self.assertEqual(contact.primary_granted, 7)
        self.assertEqual(contact.collegial_complement, 4)
        self.assertEqual(contact.qualifying_granted, 6)
        self.assertEqual(absence.collegial_granted, 5)
        self.assertEqual(absence.primary_granted, 1)
        self.assertEqual(absence.qualifying_granted, 3)

    def test_invalid_batch_is_rejected_before_any_row_is_saved(self) -> None:
        with self.assertRaises(ValueError):
            save_bulk_daily_entries([
                BulkDailyEntry("2026-06-01", MEAL_FTOUR, 90, 2),
                BulkDailyEntry("2026-06-01", MEAL_GHADA, 20, 21),
            ])

        self.assertEqual(database.get_day_contacts("2026-06-01"), [])
        self.assertEqual(database.get_day_absences("2026-06-01"), [])

    def test_template_is_rtl_and_explains_blank_absence_rule(self) -> None:
        path = self._workbook_path()
        day = QDate(2026, 6, 1).toPython()
        meal_count = create_bulk_workbook(path, day, day)

        workbook = load_workbook(path)
        self.addCleanup(workbook.close)
        sheet = workbook[SHEET_TITLE]
        self.assertEqual(meal_count, 3)
        self.assertTrue(sheet.sheet_view.rightToLeft)
        self.assertEqual([sheet.cell(5, col).value for col in range(1, 6)], [
            "التاريخ", "اليوم", "الوجبة", "الحضور", "الغياب",
        ])
        self.assertIn("فارغة", sheet["A3"].value)
        self.assertIn("0", sheet["A3"].value)
        self.assertEqual([sheet.cell(row, 6).value for row in range(6, 9)], [
            MEAL_FTOUR, MEAL_GHADA, MEAL_ASHA,
        ])
        self.assertTrue(sheet.column_dimensions["F"].hidden)
        self.assertEqual(len(sheet.data_validations.dataValidation), 1)
        self.assertEqual(workbook["_meta"].sheet_state, "hidden")

    def test_blank_absence_imports_and_saves_as_zero(self) -> None:
        path = self._workbook_path()
        day = QDate(2026, 6, 1).toPython()
        create_bulk_workbook(path, day, day)
        workbook = load_workbook(path)
        sheet = workbook[SHEET_TITLE]
        for row, attendance in zip(range(6, 9), (90, 95, 88)):
            sheet.cell(row, 4, attendance)
            sheet.cell(row, 5, None)
        workbook.save(path)
        workbook.close()

        imported = load_bulk_workbook(path)
        self.assertEqual(imported.blank_absence_count, 3)
        self.assertEqual([entry.absence for entry in imported.entries], [0, 0, 0])
        save_bulk_daily_entries(imported.entries)

        absences = database.get_day_absences("2026-06-01")
        self.assertEqual(len(absences), 3)
        self.assertTrue(all(item.collegial_granted == 0 for item in absences))

    def test_template_prefills_existing_values(self) -> None:
        save_bulk_daily_entries([
            BulkDailyEntry("2026-06-01", MEAL_FTOUR, 90, 3),
        ])
        path = self._workbook_path()
        day = QDate(2026, 6, 1).toPython()
        create_bulk_workbook(path, day, day)

        workbook = load_workbook(path)
        self.addCleanup(workbook.close)
        sheet = workbook[SHEET_TITLE]
        self.assertEqual(sheet["D6"].value, 90)
        self.assertEqual(sheet["E6"].value, 3)
        self.assertIsNone(sheet["D7"].value)
        self.assertIsNone(sheet["E7"].value)

    def test_holiday_is_not_imported_and_ramadan_uses_two_meals(self) -> None:
        database.save_school_settings(SchoolSettings(
            school_name="اختبار", school_year="2026/2027", director="مدير",
            ramadan_start="2026-03-01", ramadan_end="2026-03-30",
        ))
        database.add_holiday(Holiday(date="2026-03-02", label="عطلة"))
        path = self._workbook_path()
        create_bulk_workbook(
            path, QDate(2026, 3, 1).toPython(), QDate(2026, 3, 2).toPython(),
        )

        workbook = load_workbook(path)
        self.addCleanup(workbook.close)
        sheet = workbook[SHEET_TITLE]
        self.assertEqual([sheet.cell(row, 6).value for row in range(6, 9)], [
            MEAL_IFTAR, MEAL_SHOUR, None,
        ])
        self.assertIn("عطلة", sheet["C8"].value)
        self.assertEqual(sheet["D8"].fill.fgColor.rgb, "00E5E7EB")

    def test_import_rejects_absence_without_attendance(self) -> None:
        path = self._workbook_path()
        day = QDate(2026, 6, 1).toPython()
        create_bulk_workbook(path, day, day)
        workbook = load_workbook(path)
        workbook[SHEET_TITLE]["E6"] = 2
        workbook.save(path)
        workbook.close()

        with self.assertRaisesRegex(ValueError, "الحضور"):
            load_bulk_workbook(path)

    def test_import_rejects_absence_greater_than_attendance(self) -> None:
        path = self._workbook_path()
        day = QDate(2026, 6, 1).toPython()
        create_bulk_workbook(path, day, day)
        workbook = load_workbook(path)
        workbook[SHEET_TITLE]["D6"] = 2
        workbook[SHEET_TITLE]["E6"] = 3
        workbook.save(path)
        workbook.close()

        with self.assertRaises(ValueError):
            load_bulk_workbook(path)

    def test_dialog_displays_the_blank_absence_notice(self) -> None:
        dialog = BulkDailyEntryDialog(QDate(2026, 6, 1))
        self.addCleanup(dialog.close)
        label_text = " ".join(label.text() for label in dialog.findChildren(QLabel))
        self.assertIn("خانة الغياب اختيارية", label_text)
        self.assertIn("0", label_text)

    def test_dialog_imports_workbook_and_saves_blank_absence_as_zero(self) -> None:
        path = self._workbook_path()
        day = QDate(2026, 6, 1).toPython()
        create_bulk_workbook(path, day, day)
        workbook = load_workbook(path)
        workbook[SHEET_TITLE]["D6"] = 90
        workbook.save(path)
        workbook.close()

        dialog = BulkDailyEntryDialog(QDate(2026, 6, 1))
        self.addCleanup(dialog.close)
        with (
            patch.object(QFileDialog, "getOpenFileName", return_value=(str(path), "")),
            patch.object(
                QMessageBox, "question", return_value=QMessageBox.StandardButton.Yes,
            ),
            patch.object(
                QMessageBox, "information", return_value=QMessageBox.StandardButton.Ok,
            ),
        ):
            dialog._import_workbook()

        self.assertEqual(dialog.saved_count, 1)
        self.assertEqual(dialog.result(), QDialog.DialogCode.Accepted)
        absence = database.get_day_absences("2026-06-01")[0]
        self.assertEqual(absence.collegial_granted, 0)


if __name__ == "__main__":
    unittest.main()
