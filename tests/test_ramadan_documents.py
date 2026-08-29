"""End-to-end checks that EVERY document adapts to Ramadan.

The user's requirement was "when the user adds Ramadan mode, all the docs
adapt". Each test here drives a real document writer for a normal day and a
Ramadan day and asserts the produced file actually differs — the printed
sheets keep their own template, with two meals instead of three.
"""
import sys
import tempfile
import unittest
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path

from PySide6.QtWidgets import QApplication


ROOT_DIR = Path(__file__).resolve().parents[1]
SRC_DIR = ROOT_DIR / "src"

sys.path.insert(0, str(SRC_DIR))
sys.path.insert(0, str(ROOT_DIR))

from config.settings import (
    MEAL_ASHA, MEAL_FTOUR, MEAL_GHADA, MEAL_IFTAR, MEAL_SHOUR,
)
from core.models import (
    DailyContact, DailyReceptionRecord, SchoolSettings,
)
from data import database

_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_NS = {"w": _W}

NORMAL_DAY = "2026-01-15"
RAMADAN_DAY = "2026-03-02"


def _table_rows(path: Path, table_index: int = 0) -> list:
    with zipfile.ZipFile(path) as docx:
        root = ET.fromstring(docx.read("word/document.xml"))
    table = root.findall(".//w:tbl", _NS)[table_index]
    return [
        ["".join(t.text or "" for t in cell.findall(".//w:t", _NS)).strip()
         for cell in row.findall("./w:tc", _NS)]
        for row in table.findall("./w:tr", _NS)
    ]


class RamadanDocumentsTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.app = QApplication.instance() or QApplication([])

    def setUp(self) -> None:
        self._tmpdir = tempfile.TemporaryDirectory()
        self.out = Path(self._tmpdir.name)
        self._original_db_path = database.DB_PATH
        database.DB_PATH = self.out / "test_matama.db"
        database.init_database()
        database.save_school_settings(SchoolSettings(
            school_name="ثانوية اختبار", school_name_fr="Test College",
            school_year="2025/2026", director="MOHAMED", gestionnaire="ABDELLAH",
            company_name="STE TEST", contract_number="C/1",
            city="ألمدون", city_fr="VILLE",
            ramadan_start="2026-02-18", ramadan_end="2026-03-19",
        ))
        for meal in (MEAL_FTOUR, MEAL_GHADA, MEAL_ASHA):
            database.save_daily_contact(DailyContact(
                date=NORMAL_DAY, meal_type=meal, collegial_granted=40))
        for meal in (MEAL_IFTAR, MEAL_SHOUR):
            database.save_daily_contact(DailyContact(
                date=RAMADAN_DAY, meal_type=meal, collegial_granted=25,
                qualifying_granted=7))

    def tearDown(self) -> None:
        database.DB_PATH = self._original_db_path
        self._tmpdir.cleanup()

    # ── ورقة الاتصال اليومية ────────────────────────────────────────────────

    def test_contact_sheet_prints_two_meal_columns_in_ramadan(self) -> None:
        from ui.daily_contact_screen import _write_daily_contact_docx

        for date_str, meals, expected_cells in (
            (NORMAL_DAY, (MEAL_FTOUR, MEAL_GHADA, MEAL_ASHA), 7),
            (RAMADAN_DAY, (MEAL_IFTAR, MEAL_SHOUR), 5),
        ):
            path = self.out / f"contact_{date_str}.docx"
            _write_daily_contact_docx(
                path, date_str,
                [DailyContact(date=date_str, meal_type=m, collegial_granted=9)
                 for m in meals],
                document_number="1", place="ألمدون")
            rows = _table_rows(path)
            self.assertEqual(len(rows[3]), expected_cells, date_str)

        ramadan_rows = _table_rows(self.out / f"contact_{RAMADAN_DAY}.docx")
        self.assertEqual(ramadan_rows[1], ["", "إفطار", "سحور"])

    # ── رسالة الطلبية ──────────────────────────────────────────────────────

    def test_order_letter_lists_the_days_own_meals(self) -> None:
        from ui.order_letter_screen import (
            _order_items_from_contacts, _write_order_letter_docx,
        )
        settings = database.get_school_settings()

        for date_str, expected in (
            (NORMAL_DAY, ["فطور", "غداء", "عشاء"]),
            (RAMADAN_DAY, ["إفطار", "سحور"]),
        ):
            path = self.out / f"order_{date_str}.docx"
            items = _order_items_from_contacts(database.get_day_contacts(date_str))
            _write_order_letter_docx(path, settings, date_str, "7", items)
            rows = _table_rows(path)
            self.assertEqual([row[0] for row in rows[1:]], expected, date_str)

    # ── محضر التسلم اليومي ─────────────────────────────────────────────────

    def test_reception_record_lists_the_days_own_meals_in_french(self) -> None:
        from ui.daily_reception_screen import _write_reception_docx
        settings = database.get_school_settings()

        normal = self.out / "reception_normal.docx"
        _write_reception_docx(normal, NORMAL_DAY, DailyReceptionRecord(
            date=NORMAL_DAY, ftour_qty=11, ghada_qty=22, asha_qty=33), settings)
        ramadan = self.out / "reception_ramadan.docx"
        _write_reception_docx(ramadan, RAMADAN_DAY, DailyReceptionRecord(
            date=RAMADAN_DAY, ftour_ramadan_qty=44, shour_qty=55), settings)

        def designations(path: Path) -> list:
            with zipfile.ZipFile(path) as docx:
                root = ET.fromstring(docx.read("word/document.xml"))
            found = []
            for table in root.findall(".//w:tbl", _NS):
                for row in table.findall("./w:tr", _NS):
                    cells = ["".join(t.text or "" for t in c.findall(".//w:t", _NS)).strip()
                             for c in row.findall("./w:tc", _NS)]
                    if len(cells) >= 2 and cells[1].startswith("Le "):
                        found.append((cells[1], cells[-1]))
            return found

        self.assertEqual(
            designations(normal),
            [("Le petit-déjeuner", "11"), ("Le déjeuner", "22"), ("Le dîner", "33")])
        self.assertEqual(
            designations(ramadan), [("Le Ftour", "44"), ("Le Shour", "55")])

    # ── محضر التسلم الفصلي (the attestation the administration pays on) ────

    def test_attestation_carries_real_ramadan_counts(self) -> None:
        import openpyxl
        from ui.quarterly_reception_screen import (
            QuarterlyReceptionScreen, _write_quarterly_reception_excel,
        )
        screen = QuarterlyReceptionScreen()
        screen._year_combo.setCurrentText("2026")
        screen._first_month_combo.setCurrentIndex(0)   # Jan / Feb / Mar
        self.app.processEvents()
        # February and March both fall inside the configured Ramadan period.
        self.assertEqual(
            [m["is_ramadan"] for m in screen._quarter_data], [False, True, True])

        path = self.out / "attestation.xlsx"
        _write_quarterly_reception_excel(
            path, database.get_school_settings(), screen._quarter_data)
        book = openpyxl.load_workbook(path)

        january = book["MOIS 01 2026"]
        self.assertEqual(january.max_column, 9)        # no Ramadan blocks
        march = book["MOIS 03 2026"]
        self.assertEqual(march.max_column, 17)         # through column Q
        # March 2 = row 5. LOT 901 = primary + collegial + monitors.
        self.assertEqual(march["L5"].value, 25)        # إفطار
        self.assertEqual(march["M5"].value, 25)        # سحور
        self.assertEqual(march["P5"].value, 7)         # إفطار, qualifying
        screen.close()

    # ── بيانات المصاريف ────────────────────────────────────────────────────

    def test_expense_roster_only_shows_ramadan_columns_in_ramadan(self) -> None:
        import openpyxl
        from core.models import Student
        from ui.expense_roster_export import write_expense_roster_excel

        students = [Student(full_name="تلميذ", cycle="الثانوي الإعدادي",
                            massar_number="M1", grant_number="G1")]
        settings = database.get_school_settings()
        quarter = [(2026, 1), (2026, 2), (2026, 3)]

        without = self.out / "roster_normal.xlsx"
        write_expense_roster_excel(without, settings, students, quarter, False)
        self.assertEqual(openpyxl.load_workbook(without)["اعدادي"].max_column, 12)

        with_ramadan = self.out / "roster_ramadan.xlsx"
        write_expense_roster_excel(with_ramadan, settings, students, quarter, True)
        self.assertEqual(openpyxl.load_workbook(with_ramadan)["اعدادي"].max_column, 14)

    # ── the shared decision every document relies on ───────────────────────

    def test_every_screen_agrees_on_which_meals_a_date_serves(self) -> None:
        from ui import (
            daily_absence_screen, daily_contact_screen, daily_reception_screen,
            daily_report_screen, order_letter_screen,
        )
        modules = (daily_contact_screen, daily_absence_screen, daily_report_screen,
                   order_letter_screen, daily_reception_screen)
        for module in modules:
            normal = [key for key, _label in module._meals_for_document(NORMAL_DAY)]
            ramadan = [key for key, _label in module._meals_for_document(RAMADAN_DAY)]
            self.assertEqual(normal, [MEAL_FTOUR, MEAL_GHADA, MEAL_ASHA],
                             module.__name__)
            self.assertEqual(ramadan, [MEAL_IFTAR, MEAL_SHOUR], module.__name__)

    def test_turning_ramadan_off_restores_normal_meals_everywhere(self) -> None:
        """A school that never configures Ramadan must see the app behave
        exactly as it did before the feature existed."""
        from ui import daily_contact_screen

        settings = database.get_school_settings()
        settings.ramadan_start = ""
        settings.ramadan_end = ""
        database.save_school_settings(settings)
        self.assertEqual(
            [key for key, _l in daily_contact_screen._meals_for_document(RAMADAN_DAY)],
            [MEAL_FTOUR, MEAL_GHADA, MEAL_ASHA])


    # ── regressions from the user's own testing ────────────────────────────

    def test_report_and_order_letter_pdfs_export_without_crashing(self) -> None:
        """Both writers referenced a `date_str` that did not exist in their
        scope, so التقرير اليومي raised "name 'date_str' is not defined" and
        رسالة الطلبية produced an empty PDF (the exception aborted the draw
        half-way through the page)."""
        from ui.daily_report_screen import _report_for_date, _write_daily_report_pdf
        from ui.order_letter_screen import (
            _order_items_from_contacts, _write_order_letter_pdf,
        )
        settings = database.get_school_settings()

        for date_str in (NORMAL_DAY, RAMADAN_DAY):
            report_pdf = self.out / f"report_{date_str}.pdf"
            _write_daily_report_pdf(
                report_pdf, settings, date_str, _report_for_date(date_str))
            self.assertGreater(report_pdf.stat().st_size, 1000, date_str)

            letter_pdf = self.out / f"letter_{date_str}.pdf"
            _write_order_letter_pdf(
                letter_pdf, settings, date_str, "7",
                _order_items_from_contacts(database.get_day_contacts(date_str)))
            self.assertGreater(letter_pdf.stat().st_size, 1000, date_str)

    def test_auto_fill_produces_numbers_on_a_ramadan_day(self) -> None:
        """"تلقائي" filled nothing on a Ramadan day: the estimate builder
        only ever produced the three normal meals."""
        from ui.daily_contact_screen import _counts_to_contacts

        counts = {
            "primary": {"full": 10, "lunch": 2},
            "collegial": {"full": 20, "lunch": 3},
            "qualifying": {"full": 5, "lunch": 0},
            "monitors": {"full": 2, "lunch": 0},
        }
        ramadan = _counts_to_contacts(RAMADAN_DAY, counts)
        self.assertEqual([c.meal_type for c in ramadan], [MEAL_IFTAR, MEAL_SHOUR])
        self.assertTrue(all(c.grand_total > 0 for c in ramadan))
        # إفطار carries the وجبة غذاء students, since Ramadan has no غداء.
        iftar = next(c for c in ramadan if c.meal_type == MEAL_IFTAR)
        self.assertEqual(iftar.collegial_complement, 3)

        # A normal day must still get exactly the three normal meals.
        normal = _counts_to_contacts(NORMAL_DAY, counts)
        self.assertEqual(
            [c.meal_type for c in normal], [MEAL_FTOUR, MEAL_GHADA, MEAL_ASHA])

    def test_attestation_never_numbers_more_than_thirty_ramadan_days(self) -> None:
        """Ramadan is a Hijri month of at most 30 days, so the Ramadan block
        must not number every day of a 31-day Gregorian month."""
        import openpyxl
        from ui.quarterly_reception_screen import (
            QuarterlyReceptionScreen, _write_quarterly_reception_excel,
        )
        screen = QuarterlyReceptionScreen()
        screen._year_combo.setCurrentText("2026")
        screen._first_month_combo.setCurrentIndex(0)   # Jan / Feb / Mar
        self.app.processEvents()
        path = self.out / "attestation_days.xlsx"
        _write_quarterly_reception_excel(
            path, database.get_school_settings(), screen._quarter_data)
        book = openpyxl.load_workbook(path)

        numbered = 0
        for sheet_name, total_row in (("MOIS 02 2026", 32), ("MOIS 03 2026", 35)):
            sheet = book[sheet_name]
            days = [sheet.cell(row=row, column=11).value
                    for row in range(4, total_row)]
            numbered += sum(1 for day in days if day is not None)
        # 2026-02-18 .. 2026-03-19 inclusive is exactly 30 days.
        self.assertEqual(numbered, 30)
        self.assertLessEqual(numbered, 30)
        screen.close()


if __name__ == "__main__":
    unittest.main()
