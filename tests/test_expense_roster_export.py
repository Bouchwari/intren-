import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

import openpyxl


ROOT_DIR = Path(__file__).resolve().parents[1]
SRC_DIR = ROOT_DIR / "src"

sys.path.insert(0, str(SRC_DIR))
sys.path.insert(0, str(ROOT_DIR))

from core.models import SchoolSettings, Student
from ui import expense_roster_export as ere


def _first_student_row(sheet_name: str) -> int:
    """Students start 3 rows below that sheet's header row — and the template
    puts اعدادي's header one row lower than the other two sheets."""
    for name, _c, _m, _mon, header_row in ere._CYCLE_SHEETS:
        if name == sheet_name:
            return header_row + 3
    raise AssertionError(sheet_name)


def _student(**kwargs) -> Student:
    defaults = dict(
        full_name="تلميذ", massar_number="M1", gender="male",
        cycle="الثانوي الإعدادي", student_class="الأولى إعدادي عام",
        grant_number="G1", section="internat", grant_type="full",
    )
    defaults.update(kwargs)
    return Student(**defaults)


class ExpenseRosterExportTests(unittest.TestCase):
    """بيانات المصاريف — the official per-student quarterly roster. The
    structure asserted here was read out of the real template, not invented."""

    def setUp(self) -> None:
        self._tmpdir = tempfile.TemporaryDirectory()
        self.out = Path(self._tmpdir.name) / "roster.xlsx"
        self.settings = SchoolSettings(
            school_name="الثانوية الإعدادية ألمدون", school_year="2025-2026",
            director="MOHAMED TEST", gestionnaire="AHMED TEST",
        )
        self.quarter = [(2026, 1), (2026, 2), (2026, 3)]

    def tearDown(self) -> None:
        self._tmpdir.cleanup()

    def _write(self, students, has_ramadan: bool = False):
        ere.write_expense_roster_excel(
            self.out, self.settings, students, self.quarter, has_ramadan)
        return openpyxl.load_workbook(self.out)

    @staticmethod
    def _structure_col(has_ramadan: bool) -> int:
        """بنية الاستقبال sits right after the meal columns, so its position
        moves with them: 7 identity columns + 3 meals (+2 in Ramadan)."""
        return 8 + (5 if has_ramadan else 3)

    def test_one_sheet_per_cycle_named_like_the_template(self) -> None:
        wb = self._write([_student()])
        self.assertEqual(wb.sheetnames, ["ابتدائي", "اعدادي", "تأهيلي"])
        self.assertTrue(wb["اعدادي"].sheet_view.rightToLeft)

    def test_students_are_routed_to_their_cycle_sheet(self) -> None:
        wb = self._write([
            _student(full_name="ابن الابتدائي", cycle="الابتدائي"),
            _student(full_name="ابن الإعدادي", cycle="الثانوي الإعدادي"),
            _student(full_name="ابن التأهيلي", cycle="الثانوي التأهيلي"),
        ])
        for sheet, name in (("ابتدائي", "ابن الابتدائي"), ("اعدادي", "ابن الإعدادي"),
                            ("تأهيلي", "ابن التأهيلي")):
            self.assertEqual(
                wb[sheet].cell(row=_first_student_row(sheet), column=2).value, name)

    def test_cycle_routing_also_understands_massar_codes(self) -> None:
        """Imported rosters can carry 1A/2A/3A instead of the Arabic label."""
        wb = self._write([
            _student(full_name="بالرمز", cycle="1A"),
        ])
        self.assertEqual(
            wb["ابتدائي"].cell(row=_first_student_row("ابتدائي"), column=2).value, "بالرمز")

    def test_identity_columns_match_the_templates_order(self) -> None:
        wb = self._write([_student(
            full_name="أحمد العلوي", massar_number="M123", gender="female",
            grant_number="G9", student_class="الثانية إعدادي عام",
            grant_type="full", section="cantine")])
        ws = wb["اعدادي"]
        row = _first_student_row("اعدادي")
        self.assertEqual(ws.cell(row=row, column=1).value, 1)              # ر.ت
        self.assertEqual(ws.cell(row=row, column=2).value, "أحمد العلوي")   # الاسم
        self.assertEqual(ws.cell(row=row, column=3).value, "M123")          # رقم مسار
        self.assertEqual(ws.cell(row=row, column=4).value, "أنثى")          # الجنس
        self.assertEqual(ws.cell(row=row, column=5).value, "G9")            # رقم المنحة
        self.assertEqual(ws.cell(row=row, column=6).value, "الثانية إعدادي عام")
        self.assertEqual(ws.cell(row=row, column=7).value, "منحة كاملة")     # نوع المنحة
        self.assertEqual(
            ws.cell(row=row, column=self._structure_col(False)).value, "مطعم مدرسي")

    def test_meal_count_cells_are_left_blank_never_invented(self) -> None:
        """The standing rule for this app: the roster shows real per-student
        records, and the app has no per-student attendance data, so those
        cells stay empty for the user to fill in by hand."""
        wb = self._write([_student()], has_ramadan=True)
        ws = wb["اعدادي"]
        row = _first_student_row("اعدادي")
        for col in range(8, 13):  # the five meal columns H..L
            self.assertIsNone(ws.cell(row=row, column=col).value)

    def test_ramadan_columns_are_absent_outside_ramadan(self) -> None:
        """The user asked for السحور/الإفطار to stay OUT of the document until
        Ramadan, rather than printing two empty columns all year."""
        wb = self._write([_student()], has_ramadan=False)
        header = _first_student_row("اعدادي") - 2
        collegial = [wb["اعدادي"].cell(row=header, column=c).value for c in range(8, 11)]
        self.assertEqual(collegial, ["الفطور", "الغذاء", "العشاء"])
        self.assertIsNone(wb["اعدادي"].cell(row=header, column=11).value)
        self.assertEqual(wb["اعدادي"].max_column, 12)   # 7 identity + 3 meals + 2

    def test_ramadan_columns_appear_for_a_ramadan_quarter(self) -> None:
        wb = self._write([_student()], has_ramadan=True)
        header = _first_student_row("اعدادي") - 2
        collegial = [wb["اعدادي"].cell(row=header, column=c).value for c in range(8, 13)]
        self.assertEqual(collegial, ["الفطور", "الغذاء", "العشاء", "السحور", "الفطور"])
        self.assertEqual(wb["اعدادي"].max_column, 14)

    def test_primary_sheet_never_takes_ramadan_columns(self) -> None:
        """ابتدائي has only the three meals in the template, Ramadan or not."""
        wb = self._write([_student(cycle="الابتدائي")], has_ramadan=True)
        primary_header = _first_student_row("ابتدائي") - 2
        primary = [wb["ابتدائي"].cell(row=primary_header, column=c).value for c in range(8, 11)]
        self.assertEqual(primary, ["الفطور", "الغذاء", "العشاء"])
        self.assertIsNone(wb["ابتدائي"].cell(row=primary_header, column=11).value)

    def test_monitors_go_in_their_own_table_and_only_once(self) -> None:
        """معلمو الداخلية serve the whole internat, so listing them on more
        than one cycle sheet would claim the same people twice."""
        wb = self._write([
            _student(full_name="تلميذ عادي"),
            _student(full_name="سعيد المراقب", is_monitor=True),
            _student(full_name="ابن التأهيلي", cycle="الثانوي التأهيلي"),
        ])
        collegial_text = "\n".join(
            str(c.value) for row in wb["اعدادي"].iter_rows() for c in row if c.value)
        qualifying_text = "\n".join(
            str(c.value) for row in wb["تأهيلي"].iter_rows() for c in row if c.value)
        self.assertIn("سعيد المراقب", collegial_text)
        self.assertNotIn("سعيد المراقب", qualifying_text)
        # The monitor must not be mixed into the students table either.
        self.assertNotEqual(
            wb["اعدادي"].cell(row=_first_student_row("اعدادي"), column=2).value, "سعيد المراقب")

    def test_grant_tally_box_uses_live_countif_formulas(self) -> None:
        wb = self._write([_student(), _student(full_name="ثان", grant_type="half")])
        ws = wb["اعدادي"]
        label_col = self._structure_col(False)
        self.assertEqual(ws.cell(row=8, column=label_col).value, "منحة كاملة")
        self.assertTrue(
            str(ws.cell(row=8, column=label_col + 1).value).startswith("=COUNTIF("))
        # A half grant prints as وجبة غذاء — the wording the official document
        # uses, and the user confirmed the two mean the same thing.
        self.assertEqual(
            ws.cell(row=_first_student_row("اعدادي") + 1, column=7).value, "وجبة غذاء")

    def test_quarter_title_and_school_name_are_filled_in(self) -> None:
        wb = self._write([_student()])
        ws = wb["اعدادي"]
        self.assertIn("الثانوية الإعدادية ألمدون", str(ws["A6"].value))
        self.assertIn("الثانوي الإعدادي", str(ws["H6"].value))
        self.assertEqual(str(ws["A7"].value), "بيانات مصاريف : يناير- فبراير- مارس 2026")

    def test_year_crossing_quarter_names_the_year_on_each_month(self) -> None:
        ere.write_expense_roster_excel(
            self.out, self.settings, [_student()], [(2026, 12), (2027, 1), (2027, 2)])
        title = str(openpyxl.load_workbook(self.out)["اعدادي"]["A7"].value)
        self.assertIn("دجنبر 2026", title)
        self.assertIn("يناير 2027", title)

    def test_ministry_header_image_is_embedded_without_pillow(self) -> None:
        """Pillow is not a dependency of this project, so the crest is written
        into the package directly; the result must still be a valid workbook."""
        self._write([_student()])
        with zipfile.ZipFile(self.out) as archive:
            names = archive.namelist()
            self.assertIn("xl/media/image1.png", names)
            self.assertIn("xl/drawings/drawing1.xml", names)
            self.assertIn("xl/worksheets/_rels/sheet1.xml.rels", names)
            sheet = archive.read("xl/worksheets/sheet1.xml").decode("utf-8")
            self.assertIn("<drawing", sheet)
            content_types = archive.read("[Content_Types].xml").decode("utf-8")
            self.assertIn('Extension="png"', content_types)
        # Still loadable — the hand-written parts didn't corrupt the package.
        self.assertEqual(
            openpyxl.load_workbook(self.out).sheetnames,
            ["ابتدائي", "اعدادي", "تأهيلي"])

    def test_signature_row_matches_the_templates_spans(self) -> None:
        wb = self._write([_student()])
        text = {c.coordinate: c.value
                for row in wb["اعدادي"].iter_rows() for c in row if c.value}
        signers = [v for v in text.values() if "المدير الإقليمي" in str(v)]
        self.assertTrue(signers)
        primary_signers = "\n".join(
            str(c.value) for row in wb["ابتدائي"].iter_rows() for c in row if c.value)
        # ابتدائي has only two signers in the template.
        self.assertIn("رئيس المؤسسة", primary_signers)
        self.assertNotIn("الحارس (ة) العام للداخلية", primary_signers)

    def test_empty_roster_still_produces_a_valid_three_sheet_workbook(self) -> None:
        wb = self._write([])
        self.assertEqual(wb.sheetnames, ["ابتدائي", "اعدادي", "تأهيلي"])
        self.assertEqual(str(wb["اعدادي"]["A11"].value), "1- التلاميذ الممنوحين:")


if __name__ == "__main__":
    unittest.main()
