import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl


ROOT_DIR = Path(__file__).resolve().parents[1]
SRC_DIR = ROOT_DIR / "src"

sys.path.insert(0, str(SRC_DIR))
sys.path.insert(0, str(ROOT_DIR))

from core.excel_handler import (
    read_students_from_excel,
    write_students_template,
    write_students_to_excel,
)
from core.models import LevelOption, Student


class ExcelHandlerTests(unittest.TestCase):
    def setUp(self) -> None:
        self._tmpdir = tempfile.TemporaryDirectory()
        self.tmp_path = Path(self._tmpdir.name)

    def tearDown(self) -> None:
        self._tmpdir.cleanup()

    def test_read_students_from_excel_detects_header_variants(self) -> None:
        file_path = self.tmp_path / "students_import.xlsx"
        wb = openpyxl.Workbook()
        sheet = wb.active
        sheet.append([
            "اسم التلميذ",
            "القسم",
            "تاريخ الميلاد",
            "مكان الميلاد",
            "رقم الاستفادة",
        ])
        sheet.append(["Alice Example", "1A", "2010-01-02", "Rabat", "GR-1"])
        sheet.append(["", "ignored", "", "", ""])
        sheet.append(["Bob Example", "2B", "2011-03-04", "Casa", "GR-2"])
        wb.save(file_path)
        wb.close()

        students = read_students_from_excel(file_path)

        self.assertEqual([s.full_name for s in students], ["Alice Example", "Bob Example"])
        self.assertEqual(students[0].student_class, "1A")
        self.assertEqual(students[1].birth_place, "Casa")
        self.assertEqual(students[1].grant_number, "GR-2")

    def test_read_students_from_official_list_header_row(self) -> None:
        file_path = self.tmp_path / "official_students.xlsx"
        wb = openpyxl.Workbook()
        sheet = wb.active
        sheet.append([None, None, None, None, None, "المجموع", 1])
        sheet.append(["المؤسسة المستقبلة: مثال"])
        sheet.append([None])
        sheet.append([None])
        sheet.append([
            "ر.ت",
            "رقم التلميذ",
            "الاسم الكامل",
            "الجنس",
            "رقم المنحة",
            "المستوى",
            "بنية الاستقبال",
            "تاريخ الازدياد",
            "مكان الازدياد",
        ])
        sheet.append([
            1,
            "H194041629",
            "ايت باسو مصطفى",
            "ذكر",
            "623/25/01/391",
            "الأولى إعدادي",
            "دار الطالب/ة",
            "2010-05-15",
            "الرباط",
        ])
        wb.save(file_path)
        wb.close()

        students = read_students_from_excel(file_path)

        self.assertEqual(len(students), 1)
        self.assertEqual(students[0].full_name, "ايت باسو مصطفى")
        self.assertEqual(students[0].massar_number, "H194041629")
        self.assertEqual(students[0].gender, "male")
        self.assertEqual(students[0].section, "dar_talib")

    def test_write_students_to_excel_exports_expected_columns(self) -> None:
        file_path = self.tmp_path / "students_export.xlsx"
        write_students_to_excel(
            [
                Student(
                    full_name="Charlie Example",
                    massar_number="H0001",
                    gender="male",
                    cycle="الإعدادي",
                    education_type="عام",
                    student_class="3C",
                    birth_date="2012-05-06",
                    birth_place="Agadir",
                    grant_number="GR-3",
                    section="internat",
                    grant_type="half",
                )
            ],
            file_path,
        )

        wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        rows = list(wb.active.iter_rows(values_only=True))
        wb.close()

        self.assertEqual(rows[1], (
            1,
            "Charlie Example",
            "H0001",
            "ذكر",
            "GR-3",
            "الإعدادي",
            "عام",
            "3C",
            "القسم الداخلي",
            "2012-05-06",
            "Agadir",
            "نصف منحة",
        ))

    def test_write_students_template_creates_example_row(self) -> None:
        file_path = self.tmp_path / "students_template.xlsx"
        write_students_template(file_path)

        wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        rows = list(wb.active.iter_rows(values_only=True))
        wb.close()

        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[0][1], "الاسم الشخصي والعائلي للتلميذ(ة)")
        self.assertTrue(rows[1][0])

    def test_write_students_template_uses_filtered_level_catalog(self) -> None:
        file_path = self.tmp_path / "students_template_filtered.xlsx"
        write_students_template(
            file_path,
            [
                LevelOption("2A", "الإعدادي", "عام", "الأولى إعدادي عام"),
                LevelOption("2A", "الإعدادي", "عام", "الثانية إعدادي عام"),
            ],
        )

        wb = openpyxl.load_workbook(file_path, data_only=True)
        choices = wb["_choices"]
        levels = wb["_levels"]
        self.assertEqual(choices["A1"].value, "الإعدادي")
        self.assertEqual(choices["B1"].value, "عام")
        self.assertEqual(levels["C1"].value, "الأولى إعدادي عام")
        self.assertEqual(levels["C2"].value, "الثانية إعدادي عام")
        wb.close()


if __name__ == "__main__":
    unittest.main()
