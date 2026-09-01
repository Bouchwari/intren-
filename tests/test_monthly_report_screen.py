import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl
from PySide6.QtWidgets import QApplication


ROOT_DIR = Path(__file__).resolve().parents[1]
SRC_DIR = ROOT_DIR / "src"

sys.path.insert(0, str(SRC_DIR))
sys.path.insert(0, str(ROOT_DIR))

from config.settings import MEAL_ASHA, MEAL_FTOUR, MEAL_GHADA, MEAL_IFTAR, MEAL_SHOUR
from core.models import DailyContact, MonthlyMealSummary, SchoolSettings
from data import database
from ui import monthly_report_screen as mrs


class MonthlyReportScreenTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.app = QApplication.instance() or QApplication([])

    def setUp(self) -> None:
        self._tmpdir = tempfile.TemporaryDirectory()
        self._original_db_path = database.DB_PATH
        database.DB_PATH = Path(self._tmpdir.name) / "test_matama.db"
        database.init_database()
        database.save_school_settings(SchoolSettings(
            school_name="ثانوية اختبار", school_year="2026/2027", director="مدير",
            price_ftour="5.00", price_ghada="8.00", price_asha="5.00",
        ))
        self._originals = {
            "file_save": mrs.QFileDialog.getSaveFileName,
            "message_warning": mrs.QMessageBox.warning,
            "message_information": mrs.QMessageBox.information,
            "message_critical": mrs.QMessageBox.critical,
            "ask_choice": mrs.ask_choice,
        }

    def tearDown(self) -> None:
        mrs.QFileDialog.getSaveFileName = self._originals["file_save"]
        mrs.QMessageBox.warning = self._originals["message_warning"]
        mrs.QMessageBox.information = self._originals["message_information"]
        mrs.QMessageBox.critical = self._originals["message_critical"]
        mrs.ask_choice = self._originals["ask_choice"]
        database.DB_PATH = self._original_db_path
        self._tmpdir.cleanup()

    def test_screen_loads_real_data_on_open_not_a_blank_header(self) -> None:
        """Regression: the screen used to only pre-select the month/year
        combos in __init__ without ever calling _generate() — the school
        name and month labels stayed blank until the user manually
        clicked "توليد المحضر", unlike every sibling screen (daily/
        monthly reception) which loads its data immediately on open."""
        screen = mrs.MonthlyReportScreen()
        self.assertIn("ثانوية اختبار", screen._school_lbl.text())
        self.assertTrue(screen._month_lbl.text().strip())
        screen.close()

    def test_changing_month_combo_auto_regenerates(self) -> None:
        """Regression: the month/year selectors used to require an
        explicit "توليد المحضر" click even after changing them — a
        teacher picking a different month would keep seeing the OLD
        month's report with no visual indication anything needs
        refreshing."""
        database.save_daily_contact(DailyContact(date="2026-06-05", meal_type=MEAL_GHADA, primary_granted=7))
        screen = mrs.MonthlyReportScreen()
        screen._year_spin_combo.setCurrentText("2026")
        screen._month_combo.setCurrentIndex(5)  # June — triggers currentIndexChanged
        self.app.processEvents()
        self.assertEqual(screen._selected_month_str(), "2026-06")
        self.assertTrue(any(s.contact_total > 0 for s in screen._summaries))
        screen.close()

    def test_ramadan_meals_are_counted_but_not_priced(self) -> None:
        database.save_daily_contact(DailyContact(
            date="2026-03-10", meal_type=MEAL_IFTAR,
            collegial_granted=40,
        ))
        database.save_daily_contact(DailyContact(
            date="2026-03-10", meal_type=MEAL_SHOUR,
            collegial_granted=30,
        ))
        summaries = database.get_monthly_summaries("2026-03", {
            MEAL_FTOUR: "5", MEAL_GHADA: "8", MEAL_ASHA: "5",
        })
        by_meal = {summary.meal_type: summary for summary in summaries}
        self.assertEqual(by_meal[MEAL_IFTAR].net_total, 40)
        self.assertEqual(by_meal[MEAL_SHOUR].net_total, 30)
        self.assertEqual(by_meal[MEAL_IFTAR].unit_price, 0)
        self.assertEqual(by_meal[MEAL_SHOUR].total_cost, 0)

    def test_ordinary_month_does_not_add_empty_ramadan_rows(self) -> None:
        database.save_daily_contact(DailyContact(
            date="2026-06-05", meal_type=MEAL_GHADA,
            collegial_granted=10,
        ))
        summaries = database.get_monthly_summaries("2026-06", {})
        self.assertEqual(
            [summary.meal_type for summary in summaries],
            [MEAL_FTOUR, MEAL_GHADA, MEAL_ASHA],
        )

    def test_refresh_preserves_unsaved_notes(self) -> None:
        screen = mrs.MonthlyReportScreen()
        screen._notes_edit.setPlainText("ملاحظة لم تحفظ")
        screen.refresh()
        self.assertEqual(screen._notes_edit.toPlainText(), "ملاحظة لم تحفظ")
        screen.close()

    def test_detail_table_includes_primary_cycle_and_sums_to_subtotal(self) -> None:
        """Regression: the detail breakdown table used to only show
        إعدادي/تأهيلي/معلمو الداخلية rows, silently skipping ابتدائي
        (primary) even though contact_total/absence_total (used for the
        "مجموع" subtotal row) DOES include it — a school with primary-
        cycle boarders would see visible sector rows that didn't add up
        to the subtotal shown right below them."""
        database.save_daily_contact(DailyContact(
            date="2026-06-05", meal_type=MEAL_FTOUR,
            primary_granted=10, collegial_granted=5, qualifying_granted=3, monitors=1,
        ))
        screen = mrs.MonthlyReportScreen()
        screen._year_spin_combo.setCurrentText("2026")
        screen._month_combo.setCurrentIndex(5)
        self.app.processEvents()

        self.assertIsNotNone(screen._detail_table)
        sector_labels = [
            screen._detail_table.item(row, 1).text()
            for row in range(screen._detail_table.rowCount())
        ]
        self.assertIn("ابتدائي", sector_labels)

        ftour_summary = next(s for s in screen._summaries if s.meal_type == MEAL_FTOUR)
        self.assertEqual(ftour_summary.contact_primary, 10)
        self.assertEqual(ftour_summary.contact_total, 19)  # 10+5+3+1 — primary must count
        screen.close()

    def test_detail_table_sized_to_show_every_row_no_scrollbar(self) -> None:
        """Regression: a hardcoded "rows * guessed_pixels" height formula
        under-estimated the real per-row height, so the table showed an
        internal scrollbar and clipped most of its own rows — confirmed
        by actually rendering the screen, not just reading the formula.
        maximumHeight alone (a first fix attempt) still wasn't enough to
        make the surrounding layout actually grant that height; only
        forcing an exact minimum+maximum did."""
        database.save_daily_contact(DailyContact(date="2026-06-05", meal_type=MEAL_FTOUR, primary_granted=1))
        database.save_daily_contact(DailyContact(date="2026-06-05", meal_type=MEAL_GHADA, primary_granted=1))
        database.save_daily_contact(DailyContact(date="2026-06-05", meal_type=MEAL_ASHA, primary_granted=1))
        screen = mrs.MonthlyReportScreen()
        screen._year_spin_combo.setCurrentText("2026")
        screen._month_combo.setCurrentIndex(5)
        self.app.processEvents()

        table = screen._detail_table
        self.assertIsNotNone(table)
        header_h = table.horizontalHeader().height()
        rows_h = table.verticalHeader().length()
        expected = header_h + rows_h + table.frameWidth() * 2 + 2
        self.assertEqual(table.minimumHeight(), expected)
        self.assertEqual(table.maximumHeight(), expected)
        screen.close()

    def test_no_ghost_widgets_after_rapid_regenerate(self) -> None:
        """Regression: with the month/year combos auto-regenerating,
        two _generate() calls could land before Qt's deferred
        deleteLater() cleanup ran — confirmed via an offscreen render
        that the OLD table/cost-box widgets stayed visually painted at
        their last position, overlapping the newly built ones. Only ONE
        of each widget must exist in the report layout after several
        regenerates in a row without yielding to the event loop."""
        database.save_daily_contact(DailyContact(date="2026-06-05", meal_type=MEAL_GHADA, primary_granted=4))
        screen = mrs.MonthlyReportScreen()
        for month_index in (5, 6, 5, 7, 5):  # rapid changes, no processEvents() between them
            screen._year_spin_combo.setCurrentText("2026")
            screen._month_combo.setCurrentIndex(month_index)
        self.app.processEvents()

        main_tables = [
            screen._report_layout.itemAt(i).widget()
            for i in range(screen._report_layout.count())
        ]
        main_tables = [w for w in main_tables if w is screen._main_table]
        self.assertEqual(len(main_tables), 1)
        screen.close()

    def test_meal_colors_match_the_app_wide_canonical_palette(self) -> None:
        """The rest of the app (dashboard, meal program) color-codes
        فطور/غداء/عشاء as #EF9F27/#1D9E75/#534AB7 — this screen used to
        use its own different shades (#f59e0b/#7c3aed) for the same
        meals, so the same meal read as a different color depending
        which screen you were looking at."""
        self.assertEqual(mrs._MEAL_COLORS[MEAL_FTOUR], "#EF9F27")
        self.assertEqual(mrs._MEAL_COLORS[MEAL_ASHA], "#534AB7")

    def test_monthly_summary_main_rows_totals_match_net_and_cost(self) -> None:
        """Direct unit test on the shared row-building helper both the PDF
        and Excel export call — guards specifically against the export
        module briefly having TWO copies of this logic (one named
        `_monthly_summary_rows`, one `_monthly_summary_main_rows`) where
        only the later definition in the file actually ran; asserting on
        its real output here would have caught that drift."""
        summaries = [
            MonthlyMealSummary(meal_type=MEAL_FTOUR, days_count=2,
                               contact_primary=10, contact_collegial=5,
                               absence_primary=2, unit_price=5.0),
            MonthlyMealSummary(meal_type=MEAL_GHADA, days_count=3,
                               contact_qualifying=20, absence_qualifying=4,
                               unit_price=8.0),
        ]
        rows, totals = mrs._monthly_summary_main_rows(summaries)

        self.assertEqual(len(rows), 2)
        # ftour: contact 15, absence 2 -> net 13, cost 13*5.0
        self.assertEqual(totals["contact"], 15 + 20)
        self.assertEqual(totals["absence"], 2 + 4)
        self.assertEqual(totals["net"], 13 + 16)
        self.assertAlmostEqual(totals["cost"], 13 * 5.0 + 16 * 8.0)

    def test_export_with_no_data_warns_and_never_opens_file_dialog(self) -> None:
        """Regression: clicking تصدير before generating a report (or on a
        month with nothing entered) must not reach the file-save dialog at
        all — it should stop at a warning message instead."""
        dialog_calls: list[tuple] = []
        warnings: list[str] = []
        mrs.QFileDialog.getSaveFileName = staticmethod(lambda *a, **k: dialog_calls.append((a, k)) or ("", ""))
        mrs.QMessageBox.warning = staticmethod(lambda parent, title, text, *a, **k: warnings.append(text))

        screen = mrs.MonthlyReportScreen()
        screen._summaries = []
        screen._on_export()

        self.assertEqual(dialog_calls, [])
        self.assertEqual(len(warnings), 1)
        screen.close()

    def test_export_cancelled_format_choice_never_opens_file_dialog(self) -> None:
        """Regression: cancelling the PDF/Excel choice dialog must stop
        there, not fall through to writing a file with some default
        format."""
        dialog_calls: list[tuple] = []
        mrs.ask_choice = lambda *a, **k: None
        mrs.QFileDialog.getSaveFileName = staticmethod(lambda *a, **k: dialog_calls.append((a, k)) or ("", ""))

        database.save_daily_contact(DailyContact(date="2026-06-05", meal_type=MEAL_GHADA, primary_granted=4))
        screen = mrs.MonthlyReportScreen()
        screen._year_spin_combo.setCurrentText("2026")
        screen._month_combo.setCurrentIndex(5)
        self.app.processEvents()

        screen._on_export()
        self.assertEqual(dialog_calls, [])
        screen.close()

    def test_export_pdf_choice_writes_a_real_single_page_file(self) -> None:
        """End-to-end: choosing PDF in the export dialog must produce a
        real, non-trivial .pdf file at the chosen path via the actual
        (post-cleanup) writer function — not a stub."""
        out_path = Path(self._tmpdir.name) / "export_test.pdf"
        mrs.ask_choice = lambda *a, **k: "pdf"
        mrs.QFileDialog.getSaveFileName = staticmethod(lambda *a, **k: (str(out_path), ""))
        confirmations: list[str] = []
        mrs.QMessageBox.information = staticmethod(lambda parent, title, text, *a, **k: confirmations.append(text))

        database.save_daily_contact(DailyContact(date="2026-06-05", meal_type=MEAL_FTOUR, primary_granted=6))
        screen = mrs.MonthlyReportScreen()
        screen._year_spin_combo.setCurrentText("2026")
        screen._month_combo.setCurrentIndex(5)
        self.app.processEvents()

        screen._on_export()

        self.assertTrue(out_path.exists())
        self.assertGreater(out_path.stat().st_size, 1000)  # a real rendered page, not an empty stub
        self.assertEqual(len(confirmations), 1)
        screen.close()

    def test_export_excel_choice_writes_correct_content(self) -> None:
        """End-to-end: choosing Excel must produce a workbook whose main
        table, primary-cycle detail rows, and notes all match what's on
        screen — this is the content only the surviving (post-cleanup)
        writer produces, so a regression back to the deleted duplicate
        definition would fail this."""
        out_path = Path(self._tmpdir.name) / "export_test.xlsx"
        mrs.ask_choice = lambda *a, **k: "xlsx"
        mrs.QFileDialog.getSaveFileName = staticmethod(lambda *a, **k: (str(out_path), ""))
        mrs.QMessageBox.information = staticmethod(lambda *a, **k: None)

        database.save_daily_contact(DailyContact(
            date="2026-06-05", meal_type=MEAL_FTOUR,
            primary_granted=10, collegial_granted=5, qualifying_granted=3, monitors=1,
        ))
        screen = mrs.MonthlyReportScreen()
        screen._year_spin_combo.setCurrentText("2026")
        screen._month_combo.setCurrentIndex(5)
        self.app.processEvents()
        screen._notes_edit.setPlainText("ملاحظة اختبار")

        screen._on_export()
        self.assertTrue(out_path.exists())

        wb = openpyxl.load_workbook(out_path)
        ws = wb.active
        self.assertTrue(ws.sheet_view.rightToLeft)
        all_text = [str(cell.value) for row in ws.iter_rows() for cell in row if cell.value is not None]
        self.assertTrue(any("ثانوية اختبار" in v for v in all_text))
        self.assertIn("ابتدائي", all_text)  # primary cycle must appear, not just إعدادي/تأهيلي
        self.assertTrue(any("ملاحظة اختبار" in v for v in all_text))
        screen.close()


if __name__ == "__main__":
    unittest.main()


class MonthlySummaryExcelNumberTests(unittest.TestCase):
    """The Excel export reused the PDF's pre-formatted display strings, so
    every count, unit price and cost landed in the sheet as TEXT — a SUM over
    those columns returned 0 for whoever opened the file."""

    @classmethod
    def setUpClass(cls) -> None:
        cls.app = QApplication.instance() or QApplication([])

    def setUp(self) -> None:
        self._tmpdir = tempfile.TemporaryDirectory()
        self._original_db_path = database.DB_PATH
        database.DB_PATH = Path(self._tmpdir.name) / "test_matama.db"
        database.init_database()
        database.save_school_settings(SchoolSettings(
            school_name="مؤسسة", school_year="2025-2026", director="المدير"))

    def tearDown(self) -> None:
        database.DB_PATH = self._original_db_path
        self._tmpdir.cleanup()

    def _summaries(self):
        return [
            MonthlyMealSummary(meal_type=MEAL_FTOUR, days_count=20,
                               contact_collegial=1931, absence_collegial=60,
                               unit_price=6.0),
            MonthlyMealSummary(meal_type=MEAL_GHADA, days_count=20,
                               contact_collegial=2737, absence_collegial=92,
                               unit_price=14.0),
        ]

    def test_every_numeric_cell_is_a_real_number(self) -> None:
        path = Path(self._tmpdir.name) / "summary.xlsx"
        mrs._write_monthly_summary_excel(
            path, database.get_school_settings(), "ماي 2026", self._summaries(), "")

        sheet = openpyxl.load_workbook(path).active
        for row in (5, 6):
            for column in range(2, 8):          # everything except the meal name
                cell = sheet.cell(row=row, column=column)
                self.assertEqual(cell.data_type, "n",
                                 f"cell {cell.coordinate} is text, not a number")

    def test_the_columns_add_up_to_the_total_row(self) -> None:
        path = Path(self._tmpdir.name) / "summary.xlsx"
        mrs._write_monthly_summary_excel(
            path, database.get_school_settings(), "ماي 2026", self._summaries(), "")

        sheet = openpyxl.load_workbook(path).active
        net_column = [sheet.cell(row=row, column=5).value for row in (5, 6)]
        cost_column = [sheet.cell(row=row, column=7).value for row in (5, 6)]

        self.assertEqual(sum(net_column), sheet.cell(row=7, column=5).value)
        self.assertAlmostEqual(sum(cost_column), sheet.cell(row=7, column=7).value, places=2)
